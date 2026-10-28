"""The processing report as PDF (io.pdf_report, the content of io.report_content) and as HTML
(io.report), with the multispectral and thermal parts; the tables of a chunk as Excel workbooks
(io.tables), the Excel entries of the tables of the interface (gui.table_export), the menu actions
and the commands ``ppg report`` / ``ppg tables``."""
from __future__ import annotations

import copy
import hashlib
import os
import re
import zlib
from pathlib import Path

import numpy as np
import pytest

import pyarchinit_photogrammetry as ppg
from pyarchinit_photogrammetry.io import report_content as rc
from pyarchinit_photogrammetry.io.pdf_report import generate_pdf_report, sample_pdf
from pyarchinit_photogrammetry.io.tables import (Sheet, cell_value, export_tables, project_tables, read_xlsx,
                                                 sheet_names, write_xlsx)

pytest.importorskip("fpdf")
pytest.importorskip("openpyxl")


# ---------------------------------------------------------------------------------------- reading a PDF back
def _objects(data: bytes) -> dict:
    """``{number: (dictionary, stream or None)}`` of the objects of a PDF written by fpdf2."""
    out = {}
    for m in re.finditer(rb"(\d+) 0 obj\s*", data):
        n, start = int(m.group(1)), m.end()
        body = data[start:data.find(b"endobj", start)]
        s = body.find(b"stream")
        if s >= 0 and b"/Length" in body[:s]:
            head = body[:s]
            length = int(re.search(rb"/Length (\d+)", head).group(1))
            raw_start = start + s + len(b"stream")
            raw_start += 2 if data[raw_start:raw_start + 2] == b"\r\n" else (1 if data[raw_start:raw_start + 1] == b"\n" else 0)
            raw = data[raw_start:raw_start + length]
            out[n] = (head, zlib.decompress(raw) if b"/FlateDecode" in head else raw)
        else:
            out[n] = (body, None)
    return out


def _cmap(stream: bytes) -> dict:
    out = {}
    for block in re.findall(rb"beginbfchar(.*?)endbfchar", stream, re.S):
        for src, dst in re.findall(rb"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>", block):
            out[int(src, 16)] = bytes.fromhex(dst.decode()).decode("utf-16-be", "replace")
    for block in re.findall(rb"beginbfrange(.*?)endbfrange", stream, re.S):
        for a, b, dst in re.findall(rb"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>", block):
            for i, code in enumerate(range(int(a, 16), int(b, 16) + 1)):
                out[code] = chr(int(dst, 16) + i)
    return out


_ESC = {b"n": b"\n", b"r": b"\r", b"t": b"\t", b"b": b"\b", b"f": b"\f", b"(": b"(", b")": b")", b"\\": b"\\"}


def _unescape(s: bytes) -> bytes:
    out, i = bytearray(), 0
    while i < len(s):
        if s[i:i + 1] == b"\\" and i + 1 < len(s):
            nxt = s[i + 1:i + 2]
            if nxt in _ESC:
                out += _ESC[nxt]
                i += 2
                continue
            octal = re.match(rb"[0-7]{1,3}", s[i + 1:i + 4])
            if octal:
                out.append(int(octal.group(0), 8) & 0xFF)
                i += 1 + len(octal.group(0))
                continue
            i += 1
            continue
        out += s[i:i + 1]
        i += 1
    return bytes(out)


def pdf_text(data: bytes) -> str:
    """The text drawn on the pages (a line per string), decoded with the ToUnicode map of each
    embedded font: a light reader of what fpdf2 writes."""
    objs = _objects(data)
    fonts = {}
    for name, num in set(re.findall(rb"/(F\d+) (\d+) 0 R", data)):
        ref = re.search(rb"/ToUnicode (\d+) 0 R", objs.get(int(num), (b"", None))[0])
        if ref:
            fonts[name.decode()] = _cmap(objs[int(ref.group(1))][1] or b"")
    token = re.compile(rb"/(F\d+)\s+[\d.]+\s+Tf|\(((?:[^()\\]|\\.)*)\)\s*Tj", re.S)
    lines = []
    for _n, (_head, stream) in sorted(objs.items()):
        if stream is None or b"BT" not in stream:
            continue
        current: dict = {}
        for m in token.finditer(stream):
            if m.group(1):
                current = fonts.get(m.group(1).decode(), {})
                continue
            raw = _unescape(m.group(2))
            lines.append("".join(current.get(int.from_bytes(raw[i:i + 2], "big"), "?") for i in range(0, len(raw) - 1, 2)))
    return "\n".join(lines)


def outline(data: bytes) -> list[str]:
    """The titles of the bookmarks of a PDF."""
    out = []
    for m in re.finditer(rb"/Title\s*(\((?:[^()\\]|\\.)*\)|<[0-9A-Fa-f]+>)", data):
        v = m.group(1)
        if v.startswith(b"<"):
            raw = bytes.fromhex(v[1:-1].decode())
            out.append(raw[2:].decode("utf-16-be") if raw.startswith(b"\xfe\xff") else raw.decode("latin-1"))
        else:
            out.append(_unescape(v[1:-1]).decode("latin-1"))
    return out


def pages(data: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page(?!s)", data))


def _titles(report) -> dict[int, list[str]]:
    out: dict[int, list[str]] = {}
    for s in rc.walk(report.sections):
        out.setdefault(s.level, []).append(s.title)
    return out


def _notes(report) -> list[str]:
    return [b.text for s in rc.walk(report.sections) for b in s.blocks if isinstance(b, rc.Paragraph)]


# ---------------------------------------------------------------------------------------- a small project
def _float_tiff(path: Path, values: np.ndarray, transform=None, crs=None, descriptions=None) -> Path:
    import rasterio

    values = np.asarray(values, np.float32)
    if values.ndim == 2:
        values = values[None]
    kw = {"crs": crs, "transform": transform} if transform is not None else {}
    with rasterio.open(path, "w", driver="GTiff", height=values.shape[1], width=values.shape[2], count=values.shape[0],
                       dtype="float32", nodata=float("nan"), **kw) as dst:
        dst.write(values)
        for i, d in enumerate(descriptions or []):
            dst.set_band_description(i + 1, d)
    return path


def small_project(tmp_path: Path):
    """Two chunks with multispectral and thermal products written here - small rasters and their
    records in ``chunk.meta``, without the chain that makes them: a reflectance orthomosaic,
    NDVI and its zones; a thermal frame (a TIFF of temperatures), the temperature orthomosaic,
    its anomalies, measuring points, a spot height, a unit, a stone and a proxy."""
    import cv2
    from PIL import Image
    from rasterio.transform import from_origin

    from pyarchinit_photogrammetry.core.camera import Camera, Sensor
    from pyarchinit_photogrammetry.engine import multispectral as ms

    doc = ppg.Document()
    doc.meta["sito"] = "Saggio"
    cm = doc.add_chunk("Campo")
    ct = doc.add_chunk("Volo termico")
    doc.save(tmp_path / "campo.ppg")
    tr = from_origin(400000.0, 4500000.0, 0.05, 0.05)
    h, w = 60, 80
    for ch in (cm, ct):
        ch.crs = "EPSG:32633"
    # multispectral
    folder = cm.data_dir / "multispectral"
    folder.mkdir(parents=True)
    bands = ["Blue", "Green", "Red", "RedEdge", "NIR"]
    data = np.random.default_rng(2).uniform(0.01, 0.1, (5, h, w)).astype(np.float32)
    data[4, :, :40] += 0.4  # vegetation on the left: NIR high
    data[:, :3, :] = np.nan
    _float_tiff(folder / "reflectance.tif", data, tr, "EPSG:32633", bands)
    cm.meta["multispectral"] = {
        "reference": "Red", "bands": bands, "captures": 12, "registration": {},
        "panel": {"camera": 99, "region": [1, 1, 5, 5], "reflectance": 0.5,
                  "factors": {b: 1.0 + 0.1 * i for i, b in enumerate(bands)}},
        "registration_quality": {"0": {"NIR": {"correlation": 0.91, "on": "RedEdge"},
                                       "Blue": {"correlation": 0.8, "on": "Green"}}},
        "balance": {"model": "planes", "spread": {b: 0.03 for b in bands}, "residual": {b: 0.12 for b in bands}},
        "reflectance": "multispectral/reflectance.tif", "reflectance_bands": bands}
    ms.index_raster(cm, "NDVI")
    ms.export_zoning(cm, "NDVI", tmp_path / "zone.gpkg", classes=2, method="kmeans", min_cells=5)
    # thermal: a capture (its visible photo and thermal frame), the orthomosaic, the anomalies
    shots = tmp_path / "volo"
    shots.mkdir()
    cv2.imwrite(str(shots / "IMG_0001_W.jpg"), np.full((48, 64, 3), 120, np.uint8))
    Image.fromarray(np.full((32, 40), 21.5, np.float32), mode="F").save(shots / "IMG_0001_T.tif")
    ct.sensors[0] = Sensor(0, "camera", 64, 48)
    ct.cameras.append(Camera(0, "IMG_0001_W", str(shots / "IMG_0001_W.jpg"), 0,
                             meta={"bands": {"Visible": str(shots / "IMG_0001_W.jpg"), "LWIR": str(shots / "IMG_0001_T.tif")},
                                   "capture": "IMG_0001"}))
    tfolder = ct.data_dir / "multispectral"
    tfolder.mkdir(parents=True)
    t = np.full((h, w), 18.0, np.float32)
    t[:, 40:] = 26.0  # a warm half
    t[0, 0] = np.nan
    _float_tiff(tfolder / "thermal.tif", t, tr, "EPSG:32633", ["temperature (degrees Celsius)"])
    anomaly = np.where(np.isfinite(t), t - np.nanmean(t), np.nan)
    _float_tiff(tfolder / "thermal_anomaly.tif", anomaly, tr, "EPSG:32633")
    ct.meta["multispectral"] = {"reference": "Visible", "bands": ["LWIR", "Visible"], "captures": 1, "panel": None,
                                "registration": {}}
    ct.meta["thermal"] = {"file": "multispectral/thermal.tif", "resolution": 0.05, "params": {"emissivity": 0.97},
                          "stats": {"p2": 18.0, "p98": 26.0, "min": 18.0, "max": 26.0, "median": 22.0},
                          "offsets": {"spread": 0.42, "0": 0.42},
                          "anomaly": {"file": "multispectral/thermal_anomaly.tif", "radius_m": 2.0}}
    ct.meta["drawings"] = {
        "units": [{"wkt": "POLYGON ((400001 4499998, 400002 4499998, 400002 4499999, 400001 4499999, 400001 4499998))",
                   "sito": "Saggio", "area": "1", "us": "12", "tipo": "US", "layer": "US",
                   "attributes": {"descrizione": "strato di crollo"}}],
        "quote": [{"x": 400001.5, "y": 4499998.5, "z": 99.5, "quota_q": "q1"}],
        "points": [{"label": "P1", "x": 400001.012, "y": 4499998.988, "z": 99.512, "temp_c": 26.0, "anomaly_c": 0.4},
                   {"label": "P2", "x": 400000.5, "y": 4499998.5, "temp_c": 18.0}]}
    ct.meta["stones"] = {"thermal": ["POLYGON ((400000 4499999, 400000.2 4499999, 400000.2 4499999.2, "
                                     "400000 4499999.2, 400000 4499999))"]}
    ct.meta["proxies"] = [{"name": "US12", "us": "12", "area": "1", "volume_m3": 0.25, "area_m2": 1.0,
                           "z_top": [99.4, 99.6], "z_bottom": [99.1, 99.3], "id_media": None}]
    doc.save()
    return doc, cm, ct


# ---------------------------------------------------------------------------------------- the report
def test_pdf_report_of_the_processed_trench(processed, tmp_path):
    """The report of the processed trench: its sections in the content and as bookmarks of the
    PDF, its text readable (Unicode font), a page per part; English and Italian; the HTML page
    from the same content (still English by default)."""
    doc, ch, _ = processed
    report = rc.collect(ch, language="en")
    titles = _titles(report)
    assert titles[0] == ["Chunk: Saggio 1"]
    for t in ("Survey data", "Camera calibration", "Reference and accuracy", "Processing parameters", "Dense cloud",
              "3D model", "Digital elevation model (DEM)", "Orthomosaic"):
        assert t in titles[1], t
    assert not any("could not be read" in n for n in _notes(report)), _notes(report)
    figures = [b for s in rc.walk(report.sections) for b in s.blocks if isinstance(b, rc.Figure)]
    assert len(figures) >= 3 and all(f.width <= rc.FIGURE_SIDE and f.height <= rc.FIGURE_SIDE for f in figures)
    dem_fig = next(f for f in figures if f.colour_bar is not None)
    assert dem_fig.colour_bar.label == "Elevation (m)" and dem_fig.metres_per_pixel > 0 and dem_fig.north
    camera_fig = next(f for f in figures if f.legend is not None)
    assert [label for _c, label in camera_fig.legend.items][:2] == ["Cameras", "Control points"]
    out = generate_pdf_report(ch, tmp_path / "report_en.pdf", language="en", title="Trench", description="Test °C Ø")
    data = out.read_bytes()
    assert data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF") and pages(data) >= 5
    assert len(data) < 6_000_000  # figures reduced
    marks = outline(data)
    for t in ("Chunk: Saggio 1", "Survey data", "Reference and accuracy", "Orthomosaic"):
        assert t in marks, t
    text = pdf_text(data)
    for t in ("Trench", "Test °C Ø", "Survey data", "Mean reprojection error", "Aligned photos", "28 of 28",
              "target 0", "Control points RMSE", "EPSG:3004", "Elevation (m)"):
        assert t in text, t
    assert re.search(rf"Page 3 of\s*{pages(data)}\b", text)  # the number of pages is written at the end
    italian = generate_pdf_report(doc, tmp_path / "report_it.pdf", language="it")
    marks = outline(italian.read_bytes())
    for t in ("Dati del rilievo", "Calibrazione delle camere", "Riferimenti e accuratezza", "Ortofoto"):
        assert t in marks, t
    assert "Errore di riproiezione medio" in pdf_text(italian.read_bytes())
    html = ch.generate_report(tmp_path / "report.html")
    page = html.read_text(encoding="utf-8")
    assert "Orthomosaic" in page and "Survey data" in page and page.count("<img") >= 3
    assert "Ortofoto" in ch.generate_report(tmp_path / "rapporto.html", language="it").read_text(encoding="utf-8")


def test_local_chunks_in_metres_have_units():
    ch = ppg.Document().add_chunk()
    assert ch.crs in (None, "", "LOCAL") and not rc._metric(ch)  # arbitrary units of an unscaled block
    ch.meta["metric"] = True  # laser scans, a model imported in metres
    assert rc._metric(ch)
    ch.meta.pop("metric")
    T = np.eye(4)
    T[:3, :3] *= 27.6  # scaled by references or bars, then shifted to a local frame
    ch.transform = T
    assert rc._metric(ch)


def test_multispectral_and_thermal_parts(tmp_path):
    """A chunk with a reflectance orthomosaic, an index and its zones gets the multispectral part
    (bands, calibration, balance, false colour, the index on its ramp with its statistics, the
    zones with their areas); a chunk of a thermal camera the thermal part (decoder, parameters,
    temperatures in degrees with their ramp, anomalies) and the measuring points; Italian and
    English."""
    doc, cm, ct = small_project(tmp_path)
    report = rc.collect(doc, language="it")
    titles = _titles(report)
    assert titles[0] == ["Chunk: Campo", "Chunk: Volo termico"]
    assert {"Multispettrale", "Termico", "Punti di misura"} <= set(titles[1])
    assert {"Ortofoto di riflettanza", "Indici", "Indice NDVI", "Zonizzazione di NDVI", "Parametri della misura",
            "Ortofoto termica", "Anomalie termiche"} <= set(titles[2])
    assert not any("non è stata letta" in n for n in _notes(report)), _notes(report)
    ms_sec = next(s for s in rc.walk(report.sections) if s.title == "Multispettrale")
    kv = dict(ms_sec.blocks[0].rows)
    assert kv["Bande"] == "Blue, Green, Red, RedEdge, NIR" and kv["Banda di riferimento"] == "Red"
    assert kv["Calibrazione radiometrica"].startswith("pannello di calibrazione")
    bands = ms_sec.blocks[1]
    nir = next(r for r in bands.rows if r[0] == "NIR")
    assert nir[2] == "0,91" and nir[3] == "RedEdge" and nir[4] == "3,0"
    ndvi = next(s for s in rc.walk(report.sections) if s.title == "Indice NDVI").blocks[0]
    assert ndvi.colour_bar.label == "NDVI" and ndvi.colour_bar.low < ndvi.colour_bar.high
    zones = next(s for s in rc.walk(report.sections) if s.title == "Zonizzazione di NDVI")
    fig = next(b for b in zones.blocks if isinstance(b, rc.Figure))
    assert fig.mime == "image/png" and len(fig.legend.items) == 2
    table = next(b for b in zones.blocks if isinstance(b, rc.Table))
    assert [r[0] for r in table.rows] == ["1", "2"] and abs(sum(float(r[4].replace(",", ".")) for r in table.rows) - 100) < 0.2
    thermal = next(s for s in rc.walk(report.sections) if s.title == "Termico")
    rows = dict(thermal.blocks[0].rows)
    assert rows["Decodifica"] == "TIFF di temperature" and rows["Fotogrammi termici"] == "1"
    assert rows["Scostamenti tra i fotogrammi (deviazione standard)"] == "0,42 °C"
    params = dict(next(s for s in rc.walk(report.sections) if s.title == "Parametri della misura").blocks[0].rows)
    assert params["Emissività"] == "0,97" and params["Distanza"] == "altezza di volo di ogni fotogramma"
    ortho = next(s for s in rc.walk(report.sections) if s.title == "Ortofoto termica")
    fig = next(b for b in ortho.blocks if isinstance(b, rc.Figure))
    assert fig.colour_bar.label == "Temperatura (°C)" and (fig.colour_bar.low, fig.colour_bar.high) == (18.0, 26.0)
    stats = next(b for b in ortho.blocks if isinstance(b, rc.Table)).rows[0]
    assert stats[0] == "18,00" and stats[1] == "26,00"
    out = generate_pdf_report(doc, tmp_path / "rapporto.pdf", language="it")
    data = out.read_bytes()
    text = pdf_text(data)
    for t in ("Multispettrale", "Falso colore: R = NIR, G = Red, B = Green", "(NIR - Red) / (NIR + Red)", "Zona 1",
              "Termico", "Temperatura (°C)", "Emissività", "Punti di misura", "P1", "26,00", "ΔT (°C)"):
        assert t in text, t
    english = generate_pdf_report(ct, tmp_path / "report.pdf", language="en")
    text = pdf_text(english.read_bytes())
    for t in ("Thermal", "Parameters of the measure", "Temperature (°C)", "Thermal anomalies", "Measurement points",
              "TIFFs of temperatures"):
        assert t in text, t
    assert "Multispectral" not in outline(english.read_bytes())  # the thermal chunk only


def test_report_language_and_numbers(tmp_path):
    w = rc.Words("it")
    assert w.number(1234.5, 1) == "1.234,5" and w.coordinate(2480001.25, 2) == "2480001,25" and w.number(-0.0001, 2) == "0,00"
    assert w.length(0.0123) == "12,3 mm" and w.resolution(0.0219) == "2,19 cm/px" and w.area(12500) == "12.500 m² (1,25 ha)"
    e = rc.Words("en")
    assert e.number(1234.5, 1) == "1,234.5" and e.degrees(21.456) == "21.46 °C" and e.significant(0.000123456789) == "0.000123457"
    with pytest.raises(ValueError):
        rc.language_of("fr")
    from pyarchinit_photogrammetry.gui import i18n

    i18n.set_current("en")
    try:
        assert rc.language_of(None) == "en"  # the language of the interface
    finally:
        i18n.set_current("it")
    assert rc.nice_ticks(99.53, 100.09) == [99.6, 99.8, 100.0] and rc.nice_length(37.0) == 20.0
    cb = rc.ColourBar(rc.DEM_STOPS, 99.53, 100.09, "m")
    assert cb.tick_decimals() == 1 and cb.colours(3).shape == (3, 3)
    assert sample_pdf(tmp_path / "sample.pdf").read_bytes().startswith(b"%PDF")


def test_pdf_without_a_unicode_font(tmp_path, monkeypatch):
    """Without DejaVu Sans (no matplotlib, no system font) the report falls back to the core font
    of PDF: the texts outside Windows-1252 (Delta, superscripts) become plain letters."""
    from pyarchinit_photogrammetry.io import pdf_report

    doc, cm, ct = small_project(tmp_path)
    monkeypatch.setattr(pdf_report, "font_files", lambda: {})
    out = pdf_report.generate_pdf_report(ct, tmp_path / "core.pdf", language="it")
    data = out.read_bytes()
    assert data.startswith(b"%PDF") and b"/Helvetica" in data and b"DejaVu" not in data
    assert {"Termico", "Punti di misura"} <= set(outline(data))


# ---------------------------------------------------------------------------------------- GNSS
DATA = Path(__file__).parent / "data"
ATX_EXCERPT = DATA / "gnss" / "igs20_excerpt.atx"  # TRM29659.00 NONE (synthetic), LEIAR25.R3 NONE and LEIT, ...
RINEX_BASE = DATA / "rtklib" / "07590920.05o"  # a RINEX header with the antenna TRM29659.00 (RTKLIB's sample data)
PLACE = (42.5, 12.5, 300.0)  # a made-up place: latitude, longitude (degrees), height above the ellipsoid (m)
EPOCH = 2026.82
EPN_CITATION = ("EUREF Permanent GNSS Network (EPN), multi-year position and velocity solution C2400 in IGS20, EPN "
                "Central Bureau, Royal Observatory of Belgium, doi:10.24414/ROB-EUREF-C2400, CC BY 4.0")
RING_CITATION = ("INGV RING, GipsyX daily solutions in IGS20 (https://doi.org/10.13127/gnss/gipos), CC BY 4.0: a "
                 "straight line through two years of daily positions")


def _station(networks=("EPN",), station_id="PROV00ITA", name="Stazione di prova"):
    from pyarchinit_photogrammetry.gnss import rinex, stations

    lat, lon, h = PLACE
    return stations.Station(station_id, name, tuple(networks), lat, lon, h, tuple(rinex.llh_to_ecef(PLACE)))


def _station_position(monkeypatch, station, source="EPN", citation=EPN_CITATION, sigma=(0.0011, 0.0009, 0.0021)):
    """The position of ``station`` as the base of a flight, as gnss.base records it: its coordinates in ITRF2020 at
    EPOCH from ``source`` with its ``citation``, given instead of downloaded."""
    from pyarchinit_photogrammetry.gnss import base, stations

    def coordinates(st, epoch, progress=None):
        return stations.StationCoordinates(st, st.xyz, (st.lat, st.lon, st.h), "ITRF2020", epoch, source, sigma,
                                           citation)

    monkeypatch.setattr(stations, "coordinates", coordinates)
    return base._station_position(station, EPOCH, None)


def _base_record(position, atx_file, name, antenna=None) -> dict:
    """``Solution.base`` as ppk.solve makes it: the base at ``position`` with the antenna of a RINEX header (or
    ``antenna``) and what the antenna file ``atx_file`` gives it (None: no file given to RTKLIB)."""
    from pyarchinit_photogrammetry.gnss import ppk, rinex

    record, _warning, _llh = ppk._base(RINEX_BASE, rinex.header(RINEX_BASE), position)
    record["name"] = name
    if antenna is not None:
        record["antenna"] = antenna
    record["antenna_calibration"] = None if atx_file is None else ppk.antenna_calibration(atx_file, record["antenna"])
    return record


def _ppk_photos(ch, record: dict, solution: str, first: int, count: int, when: str, states=("fix",)) -> None:
    """``count`` photos taken a second apart from ``when`` with the PPK positions apply_ppk gives them against the base
    ``record`` (the states of ``states`` in turn), 1.2 km north of the base."""
    import copy
    from datetime import datetime, timedelta

    from pyarchinit_photogrammetry.core.camera import Camera
    from pyarchinit_photogrammetry.gnss import apply, frames

    lon0, lat0, h0 = record["position"]
    frame = record.get("frame") or frames.UNKNOWN
    for i in range(count):
        taken = datetime.strptime(when, "%Y:%m:%d %H:%M:%S") + timedelta(seconds=i)
        cam = Camera(first + i, f"DJI_{first + i:04d}", f"DJI_{first + i:04d}.JPG", 0,
                     meta={"datetime": f"{taken:%Y:%m:%d %H:%M:%S}"})
        state = states[i % len(states)]
        lon, lat, h = lon0 + 0.0001 * i, lat0 + 0.0108, h0 + 60.0
        extra = {"mrk": "VOLO_Timestamp.MRK", "shot": i + 1, "utc": taken.isoformat(), "lever_arm_ned_mm": [0, 0, 120],
                 "antenna": [lon, lat, h + 0.12], "q": {"fix": 1, "float": 2}.get(state, 5),
                 "base": copy.deepcopy(record), "rover": "VOLO_Rinex.obs", "solution": solution}
        apply.set_camera_position(cam, [lon, lat, h], frames.crs_of(frame), apply.accuracy(state, [0.008, 0.009, 0.015]),
                                  "ppk", status=state, height="ellipsoidal", extra=extra, frame=frame,
                                  epoch=record["epoch"] if frame in frames.DYNAMIC else frames.decimal_year(taken))
        ch.cameras.append(cam)


def _exif_photos(ch, first: int, count: int, when: str) -> None:
    """``count`` photos with the GPS of their EXIF, as Chunk.add_photos gives it (no RTK data)."""
    from pyarchinit_photogrammetry.core.camera import Camera
    from pyarchinit_photogrammetry.gnss import apply

    lat, lon, h = PLACE
    for i in range(count):
        cam = Camera(first + i, f"IMG_{first + i:04d}", f"IMG_{first + i:04d}.JPG", 0, meta={"datetime": when})
        apply.set_camera_position(cam, [lon + 0.001 * i, lat + 0.011, h + 15.0], "EPSG:4326", list(apply.NO_RTK), "exif")
        ch.cameras.append(cam)


def gnss_project(tmp_path, monkeypatch):
    """A chunk whose first flight has six photos placed by PPK against a permanent station of the EPN as the base (five
    fix, one float; the station's coordinates in ITRF2020 at the epoch of the flight, its antenna calibrated by an
    excerpt of igs20.atx), and a second flight, an hour later, two photos with the GPS of their EXIF; the system of the
    chunk has heights above the geoid EGM2008."""
    import shutil

    from pyarchinit_photogrammetry.core.camera import Sensor

    doc = ppg.Document()
    ch = doc.add_chunk("Volo RTK")
    doc.save(tmp_path / "rilievo.ppg")
    ch.sensors[0] = Sensor(0, "FC6310R", 5472, 3648)
    atx_file = Path(shutil.copyfile(ATX_EXCERPT, tmp_path / "igs20.atx"))
    station = _station()
    record = _base_record(_station_position(monkeypatch, station), atx_file, station.id)
    _ppk_photos(ch, record, "solution-20261027-101500.pos", 0, 6, "2026:10:27 10:15:00", ("fix",) * 5 + ("float",))
    _exif_photos(ch, 6, 2, "2026:10:27 11:30:00")
    ch.crs = "EPSG:7792+3855"
    doc.save()
    return doc, ch


def _gnss_parts(report) -> tuple[rc.Section, dict]:
    """The GNSS part of the first chunk of ``report`` and its parts by title."""
    sec = next(s for s in report.sections[0].blocks
               if isinstance(s, rc.Section) and s.title in ("GNSS positions", "Posizioni GNSS"))
    return sec, {s.title: s for s in rc.walk([sec]) if s is not sec}


def _rows(section: rc.Section) -> dict:
    return dict(next(b for b in section.blocks if isinstance(b, rc.KeyValues)).rows)


def test_report_gnss_section(tmp_path, monkeypatch):
    """The GNSS part of the report of a chunk whose photos have PPK positions against a permanent station of the EPN:
    where the positions come from, their frames and epochs, the quality of every flight; the base of the PPK with its
    station, where the station's coordinates come from and what their sigma is, its frame and epoch, the calibration of
    its antenna named by its entry in igs20.atx (never as a number of centimetres: the L1 offset is not what moves the
    heights); the system of the chunk with the grid of its geoid and its SHA-256, the Helmert that takes ITRF2020 there;
    the credits of the station's coordinates with the DOI of the EPN; English and Italian, and in the PDF."""
    from pyarchinit_photogrammetry.gnss import heights

    doc, ch = gnss_project(tmp_path, monkeypatch)
    english = rc.collect(ch, language="en")
    assert "GNSS positions" in _titles(english)[1]
    assert not any("could not be read" in n for n in _notes(english)), _notes(english)
    sec, parts = _gnss_parts(english)
    rows = _rows(sec)
    assert rows["Photos with a position"] == "8 of 8"
    assert rows["Where the positions come from"] == "PPK: 6; EXIF GPS: 2"
    assert rows["Frame of the positions"] == "ITRF2020 at the epoch 2026.82: 6; unknown (WGS 84): 2"
    assert rows["System of the chunk"].startswith("EPSG:7792+3855 - ")
    assert rows["Heights of the chunk"] == "above the geoid (EGM2008 height)"
    assert heights.GEOID_GRID in rows["Geoid grid"] and heights.GEOID_SHA256 in rows["Geoid grid"]
    assert "ITRF2020 to ETRF2000 (1) (EPSG:10586)" in rows["Conversion of the positions"]
    flights = next(b for b in parts["Quality per flight"].blocks if isinstance(b, rc.Table))
    assert flights.headers[:7] == ["Flight", "Photos", "Fix", "Float", "Single", "No RTK", "No position"]
    assert [r[:7] for r in flights.rows] == [["1", "6", "5", "1", "0", "0", "0"], ["2", "2", "0", "0", "0", "2", "0"]]
    assert flights.rows[0][7] == "1.6 / 1.8 / 3.0 cm" and flights.rows[1][7] == "5.00 / 5.00 / 10.00 m"
    run = _rows(parts["PPK against the base PROV00ITA"])
    assert run["Solution"] == "solution-20261027-101500.pos" and run["Photos"] == "6 (fix 5, float 1)"
    assert run["Position of the base"] == "a permanent station as the base"
    assert run["Permanent station"] == "PROV00ITA - Stazione di prova (EPN)"
    assert run["Coordinates of the station"] == "EPN, ITRF2020 at the epoch 2026.82"
    assert run["Standard deviations of the station (E / N / U)"].endswith(" mm, formal error of the solution")
    assert run["Frame and epoch of the base"] == "ITRF2020 at the epoch 2026.82"
    assert run["Calibration of the base antenna"] == "calibrated with TRM29659.00 NONE of igs20.atx"
    assert not any("0.100" in v or "100.0 mm" in v or "10.0 cm" in v for v in run.values())  # never the L1 offset
    assert run["Distance from the flight (median)"] == "1.20 km"
    credits = _rows(parts["Credits"])
    assert credits["Coordinates of the station PROV00ITA"] == EPN_CITATION
    assert credits["Observations of the station PROV00ITA"] == "EUREF Permanent GNSS Network (EPN), CC BY 4.0"
    assert "igs20.atx" in credits["Antenna calibrations"] and "EGM2008" in credits["Geoid"]
    italian = rc.collect(doc, language="it")
    sec, parts = _gnss_parts(italian)
    rows = _rows(sec)
    assert rows["Foto con una posizione"] == "8 di 8"
    assert rows["Sistema delle posizioni"] == "ITRF2020 all'epoca 2026,82: 6; sconosciuto (WGS 84): 2"
    assert heights.GEOID_SHA256 in rows["Griglia del geoide"]
    run = _rows(parts["PPK rispetto alla base PROV00ITA"])
    assert run["Stazione permanente"] == "PROV00ITA - Stazione di prova (EPN)"
    assert run["Coordinate della stazione"] == "EPN, ITRF2020 all'epoca 2026,82"
    assert run["Deviazioni standard della stazione (E / N / U)"].endswith(" mm, errore formale della soluzione")
    assert run["Sistema ed epoca della base"] == "ITRF2020 all'epoca 2026,82"
    assert run["Calibrazione dell'antenna della base"] == "calibrata con TRM29659.00 NONE di igs20.atx"
    assert _rows(parts["Crediti"])["Coordinate della stazione PROV00ITA"] == EPN_CITATION
    data = generate_pdf_report(doc, tmp_path / "rapporto.pdf", language="it").read_bytes()
    assert {"Posizioni GNSS", "Qualità per volo", "PPK rispetto alla base PROV00ITA", "Crediti"} <= set(outline(data))
    text = pdf_text(data)
    for t in ("PROV00ITA - Stazione di prova (EPN)", "ITRF2020", "2026,82", "igs20.atx", "doi:10.24414/ROB-EUREF-C2400",
              heights.GEOID_GRID):
        assert t in text, t
    page = ch.generate_report(tmp_path / "gnss.html", language="en").read_text(encoding="utf-8")
    assert "GNSS positions" in page and "PPK against the base PROV00ITA" in page and "ROB-EUREF-C2400" in page


def test_report_gnss_says_itrf2020_and_a_european_frame_together_on_wgs84(tmp_path, monkeypatch):
    """The chunk of the PPK against a permanent station (ITRF2020) in WGS 84 / UTM with markers read from a file in
    RDN2008 (a network RTK): a note says they are about a metre apart there, and the system of ETRF2000 to choose."""
    from pyarchinit_photogrammetry.core import crs as crs_mod
    from pyarchinit_photogrammetry.markers import gcp

    doc, ch = gnss_project(tmp_path, monkeypatch)
    ch.crs = "EPSG:32633"
    lon, lat = (float(v) for v in np.median([c.reference.location[:2] for c in ch.cameras[:6]], axis=0))
    x, y, _z = crs_mod.transform_points([[lon, lat, 0.0]], "EPSG:4979", "EPSG:7792")[0]
    table = tmp_path / "gcp.csv"
    table.write_text(f"GCP1,{x},{y},52.0\n")
    assert gcp.import_reference(ch, table, crs="EPSG:7792") == 1
    notes = _notes(rc.collect(ch, language="en"))
    said = [n for n in notes if "about a metre apart" in n]
    assert len(said) == 1 and "ITRF2020 positions of 6 photos and 1 references" in said[0]
    assert "EPSG:32633" in said[0] and "(EPSG:7792)" in said[0]
    assert any("a circa un metro fra loro" in n for n in _notes(rc.collect(ch, language="it")))
    ch.crs = "EPSG:7792+3855"
    assert not any("about a metre apart" in n for n in _notes(rc.collect(ch, language="en")))


def test_report_gnss_bases_placed_in_other_ways(tmp_path, monkeypatch):
    """Three PPK runs in one chunk, each with its base: one placed by a static baseline from a station of RING (whose
    sigma is the scatter of its daily positions), one by a static PPP (its products named and their centre credited;
    its antenna not in the antenna file: not calibrated), one at the position of its RINEX header with no antenna
    calibrations given; the system of the chunk with heights above a geoid of the user's own (ITALGEO2005), whose grid
    is named."""
    import shutil

    from pyarchinit_photogrammetry.core.camera import Sensor
    from pyarchinit_photogrammetry.gnss import ppk, ppp

    doc = ppg.Document()
    ch = doc.add_chunk("Tre voli")
    doc.save(tmp_path / "voli.ppg")
    ch.sensors[0] = Sensor(0, "FC6310R", 5472, 3648)
    atx_file = Path(shutil.copyfile(ATX_EXCERPT, tmp_path / "igs20.atx"))
    ring = _station(("RING",), "PROV01ITA", "Stazione RING di prova")
    sp = _station_position(monkeypatch, ring, "RING", RING_CITATION, (0.0015, 0.0012, 0.0040))
    lat, lon, h = sp.lat + 0.03, sp.lon + 0.02, sp.h - 20.0
    detail = {"station": ring.id, "name": ring.name, "networks": list(ring.networks),
              "station_position": {"lat": sp.lat, "lon": sp.lon, "h": sp.h, "source": sp.source, "frame": sp.frame,
                                   "epoch": sp.epoch, "sd": list(sp.sd), "detail": sp.detail},
              "citation": sp.detail.get("citation"), "base": "BASE1", "antenna": "LEIAR25.R3 LEIT",
              "antenna_delta_hen": [1.5, 0.0, 0.0], "antenna_calibration": ppk.antenna_calibration(atx_file, "LEIAR25.R3 LEIT"),
              "baseline_km": 4.2, "q": 1, "status": "fix", "hours": 2.0, "solution": "solution-20261027-070000.pos"}
    by_baseline = ppk.BasePosition(lat, lon, h, "baseline", frame=sp.frame, epoch=sp.epoch, sd=(0.004, 0.005, 0.011),
                                   detail=detail)
    _ppk_photos(ch, _base_record(by_baseline, atx_file, "BASE1", "LEIAR25.R3 LEIT"), "solution-20261027-101500.pos", 0,
                3, "2026:10:27 10:15:00")
    products = ppp.Products([tmp_path / "ESA0OPSFIN_20263000000_01D_05M_ORB.SP3.gz"],
                            [tmp_path / "ESA0OPSFIN_20263000000_01D_30S_CLK.CLK.gz"],
                            [tmp_path / "ESA0OPSFIN_20263000000_01D_01D_ERP.ERP.gz"], "ESA", "final")
    detail = {"base": "BASE2", "centre": "ESA", "kind": "final",
              "products": [p.name for p in (*products.sp3, *products.clk, *products.erp)],
              "citation": ppp._citation(products, atx_file), "hours": 4.0,
              "session": ["2026-10-27T06:00:00", "2026-10-27T10:00:00"], "antenna": "JAVRINGANT_DM NONE",
              "antenna_delta_hen": [0.0, 0.0, 0.0], "antenna_calibration": None, "q": 6, "quality": "ppp",
              "solution": "solution-20261027-100500.pos"}
    by_ppp = ppk.BasePosition(lat + 0.05, lon, h, "ppp", frame="ITRF2020", epoch=EPOCH, sd=(0.03, 0.03, 0.06), detail=detail)
    _ppk_photos(ch, _base_record(by_ppp, atx_file, "BASE2", "JAVRINGANT_DM NONE"), "solution-20261027-130000.pos", 3, 2,
                "2026:10:27 13:00:00")
    header = ppk.BasePosition(lat - 0.04, lon, h, "header")
    _ppk_photos(ch, _base_record(header, None, "BASE3"), "solution-20261027-160000.pos", 5, 2, "2026:10:27 16:00:00")
    ch.crs = "EPSG:7792+5214"  # RDN2008 / UTM 33N and the heights of Genova 1942: ITALGEO2005, a grid of the user's own
    doc.save()
    sec, parts = _gnss_parts(rc.collect(ch, language="en"))
    rows = _rows(sec)
    assert rows["Heights of the chunk"].startswith("above the geoid (") and "geo_igm_mar06.grd" in rows["Geoid grid"]
    first = _rows(parts["PPK against the base BASE1"])
    assert first["Position of the base"] == "computed by a static baseline from a permanent station"
    assert first["Permanent station"] == "PROV01ITA - Stazione RING di prova (RING)"
    assert first["Coordinates of the station"] == "RING, ITRF2020 at the epoch 2026.82"
    assert first["Standard deviations of the station (E / N / U)"].endswith(" mm, scatter of the daily positions")
    assert first["Static baseline"] == "4.200 km from the station, fix (Q=1), 2.00 hours observed together"
    assert first["Standard deviations of the base (E / N / U)"] == "0.4 / 0.5 / 1.1 cm"
    assert first["Antenna of the base"] == "LEIAR25.R3 LEIT, 0.000 m above the mark (RINEX header)"
    assert first["Calibration of the base antenna"] == "calibrated with LEIAR25.R3 LEIT of igs20.atx"
    second = _rows(parts["PPK against the base BASE2"])
    assert second["Position of the base"] == "computed by static PPP on precise products"
    assert second["Static PPP"] == "4.00 hours, final products of ESA"
    assert second["Precise products"].startswith("ESA0OPSFIN_20263000000_01D_05M_ORB.SP3.gz, ")
    assert second["Calibration of the base antenna"] == ("not calibrated: igs20.atx has no entry for the antenna "
                                                          "JAVRINGANT_DM NONE")
    third = _rows(parts["PPK against the base BASE3"])
    assert third["Position of the base"].startswith("from the RINEX header of the base: it can be metres off")
    assert third["Frame and epoch of the base"] == "unknown (WGS 84)"
    assert third["Calibration of the base antenna"] == "not calibrated: RTKLIB was given no file of antenna calibrations"
    credits = _rows(parts["Credits"])
    assert credits["Coordinates of the station PROV01ITA"] == RING_CITATION
    assert credits["Observations of the station PROV01ITA"] == "INGV RING, CC BY 4.0"
    assert credits["Precise products"] == ppp._citation(products, atx_file)
    assert "ESA final orbits" in credits["Precise products"] and "Geoid" not in credits


def test_report_without_gnss_has_no_section(tmp_path):
    """Photos without a position, or no photos: no GNSS part."""
    from pyarchinit_photogrammetry.core.camera import Camera, Sensor

    doc = ppg.Document()
    empty = doc.add_chunk("Vuoto")
    photos = doc.add_chunk("Senza posizioni")
    photos.sensors[0] = Sensor(0, "camera", 64, 48)
    photos.cameras.append(Camera(0, "IMG_0001", "IMG_0001.JPG", 0, meta={"datetime": "2026:10:27 10:00:00"}))
    doc.save(tmp_path / "senza.ppg")
    w = rc.Words("en")
    assert rc._gnss(empty, w) is None and rc._gnss(photos, w) is None
    titles = [s.title for s in rc.walk(rc.collect(doc, language="en").sections)]
    assert "Chunk: Senza posizioni" in titles and "GNSS positions" not in titles


def test_report_hashes_the_egm2008_grid_in_use(tmp_path, monkeypatch):
    """The row of the geoid grid gives the SHA-256 of the file in use (the copy of the program), and the published one,
    so labelled, when the grid is not installed on this computer."""
    from pyarchinit_photogrammetry.gnss import heights

    monkeypatch.setenv("PPG_GNSS_DIR", str(tmp_path / "gnss"))
    w = rc.Words("en")
    (row,) = rc._geoid_grids(w, "EPSG:32633+3855")
    assert row == f"{heights.GEOID_GRID}, published SHA-256 {heights.GEOID_SHA256}, not installed on this computer"
    copy = heights.grid_folder() / heights.GEOID_GRID
    copy.parent.mkdir(parents=True)
    copy.write_bytes(b"a copy of the grid that is not the published one")
    (row,) = rc._geoid_grids(w, "EPSG:32633+3855")
    assert row == f"{heights.GEOID_GRID}, SHA-256 {hashlib.sha256(copy.read_bytes()).hexdigest()}"


def test_report_titles_runs_against_one_base_apart(tmp_path, monkeypatch):
    """Two flights against one base: two runs, told apart in their titles by the time of their first shot; a run whose
    solution was read from a file, with no base, is titled by its solution file."""
    doc, ch = gnss_project(tmp_path, monkeypatch)
    record = copy.deepcopy((ch.cameras[0].meta["gnss"])["base"])
    _ppk_photos(ch, record, "solution-20261027-121500.pos", 20, 3, "2026:10:27 12:15:00", ("fix",) * 3)
    titles = [s.title for s in rc.walk(rc.collect(ch, language="en").sections)]
    runs = [t for t in titles if t.startswith("PPK against the base")]
    assert runs == ["PPK against the base PROV00ITA, 2026-10-27 10:15", "PPK against the base PROV00ITA, 2026-10-27 12:15"]
    for cam in (c for c in ch.cameras if c.key >= 20):  # the photos of the second run
        cam.meta["gnss"]["base"] = {}
    titles = [s.title for s in rc.walk(rc.collect(ch, language="en").sections)]
    assert "PPK against the base PROV00ITA" in titles and "PPK: solution-20261027-121500.pos" in titles


def test_report_says_an_entry_without_the_offsets_of_l1_is_no_calibration():
    """An antenna whose entry in the antenna file has no offsets of L1 (G01) is said not calibrated, naming the entry."""
    base = {"antenna": "ASH700829.2     SNOW", "antenna_calibration": {
        "file": "igs20.atx", "entry": None, "without_radome": False, "up_m": None, "no_l1": "ASH700829.2     SNOW"}}
    assert rc._calibration_text(rc.Words("en"), base) == ("not calibrated: the entry ASH700829.2 SNOW of igs20.atx has "
                                                          "no offsets of L1 (G01)")
    assert rc._calibration_text(rc.Words("it"), base).startswith("non calibrata: la voce ASH700829.2 SNOW di igs20.atx")


# ---------------------------------------------------------------------------------------- Excel
def test_write_xlsx_keeps_numbers_and_makes_valid_sheets(tmp_path):
    from openpyxl import load_workbook

    names = sheet_names(["Camere", "a/b:c*d?[e]", "x" * 40, "camere", "History", "", "'quoted'"])
    assert names[0] == "Camere" and names[1] == "a_b_c_d__e_" and len(names[2]) == 31
    assert names[3] == "camere (2)" and names[4] == "History_" and names[5] == "Sheet" and names[6] == "quoted"
    assert cell_value(np.float32(1.5)) == 1.5 and cell_value(np.int64(3)) == 3 and cell_value(float("nan")) is None
    assert cell_value(True) is True and cell_value([1, 2]) == "1, 2"
    tables = {"Misure": Sheet(["Nome", "X", "Punti", "Errore", "Note", "Ok"],
                              [["P1", 2480001.123456, 12, 0.00123, "=1+1", True],
                               ["P2", np.float64(2480002.5), np.int32(40000), float("nan"), None, False]],
                              [None, 3, None, 4, None, None]),
              "Misure ": (["a"], [[1]]),
              "Progetto": (["Campo", "Valore"], [["Foto", 57], ["Punti", 316195], ["Errore", 0.7123]]),
              "Prodotti": Sheet(["Punti"], [[316195], [12]], [0])}
    out = write_xlsx(tmp_path / "t", tables, title="Prova")
    assert out.suffix == ".xlsx"
    wb = load_workbook(out)
    assert wb.sheetnames == ["Misure", "Misure (2)", "Progetto", "Prodotti"] and wb.properties.title == "Prova"
    assert wb["Prodotti"]["A2"].number_format == "#,##0"  # no decimals asked: thousands still grouped
    kv = wb["Progetto"]  # counts among measures stay whole numbers
    assert (kv["B2"].number_format, kv["B3"].number_format, kv["B4"].number_format) == ("0", "#,##0", "0.000")
    ws = wb["Misure"]
    assert [c.value for c in ws[1]] == ["Nome", "X", "Punti", "Errore", "Note", "Ok"]
    assert ws["B2"].value == pytest.approx(2480001.123456) and ws["B2"].number_format == "0.000"
    assert ws["C3"].value == 40000 and ws["C2"].number_format == "#,##0" and ws["D2"].number_format == "0.0000"
    assert ws["D3"].value is None and ws["E2"].value == "=1+1" and ws["E2"].data_type == "s"
    assert ws["F2"].value is True and ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:F3"
    assert ws["A1"].font.bold and ws.column_dimensions["A"].width >= 6
    with pytest.raises(ValueError):
        write_xlsx(tmp_path / "empty.xlsx", {})


def test_tables_of_the_processed_trench(processed, tmp_path):
    """Every table of the processed trench in a workbook: cameras with positions, angles and tie
    points, sensors, markers with their errors, products, classes, parameters - numbers as numbers."""
    _doc, ch, _ = processed
    out = export_tables(ch, tmp_path / "tabelle.xlsx", language="en")
    book = read_xlsx(out)
    for name in ("Project", "Cameras", "Sensors", "Markers", "Parameters", "Products", "Point classes", "Versions"):
        assert name in book, name
    cams = book["Cameras"]
    head = cams[0]
    assert head[:4] == ["Label", "Enabled", "Aligned", "Sensor"] and len(cams) == 1 + len(ch.cameras)
    x, tie = head.index("X estimated"), head.index("Tie points")
    assert all(isinstance(r[x], float) and r[x] > 2_479_000 for r in cams[1:])
    assert all(isinstance(r[tie], int) and r[tie] > 0 for r in cams[1:])
    markers = book["Markers"]
    err = markers[0].index("Error (m)")
    assert all(isinstance(r[err], float) and r[err] < 0.01 for r in markers[1:])
    sensors = book["Sensors"]
    assert sensors[1][sensors[0].index("Model")] == ch.sensors[0].model
    assert sensors[1][sensors[0].index("cx")] == pytest.approx(ch.sensors[0].K[0, 2])  # a number (600 reads as int)
    products = {r[0]: r for r in book["Products"][1:]}
    assert {"point_cloud", "model", "dem", "orthomosaic"} <= set(products)
    assert isinstance(products["point_cloud"][3], int) and products["point_cloud"][3] > 50_000
    italian = project_tables(ch, language="it")
    assert {"Progetto", "Camere", "Sensori", "Marker", "Parametri"} <= set(italian)


def test_tables_of_multispectral_and_thermal_products(tmp_path):
    doc, cm, ct = small_project(tmp_path)
    book = read_xlsx(export_tables(cm, tmp_path / "campo.xlsx", language="it"))
    assert {"Bande", "Indici", "Zone NDVI"} <= set(book)
    bands = {r[0]: r for r in book["Bande"][1:]}
    assert bands["Red"][1] is True and bands["NIR"][3] == pytest.approx(0.91) and isinstance(bands["NIR"][7], float)
    ndvi = book["Indici"][1]
    assert ndvi[0] == "NDVI" and ndvi[1] == "(NIR - Red) / (NIR + Red)" and -1 <= ndvi[3] < ndvi[4] <= 1
    zones = book["Zone NDVI"]
    assert [r[0] for r in zones[1:]] == [1, 2] and zones[1][1] == pytest.approx(ndvi[3], abs=1e-6)
    book = read_xlsx(export_tables(ct, tmp_path / "termico.xlsx", language="en"))
    assert {"Thermal", "Thermal offsets", "Measurement points", "Spot heights", "Units", "Stones",
            "Unit proxies"} <= set(book)
    thermal = dict((r[0], r[1]) for r in book["Thermal"][1:])
    assert thermal["Emissivity"] == 0.97 and thermal["Thermal frames"] == 1 and thermal["Decoder"] == "tiff"
    assert thermal["Minimum temperature (°C)"] == 18.0 and thermal["Maximum temperature (°C)"] == 26.0
    points = book["Measurement points"]
    assert points[0][:5] == ["Point", "X", "Y", "Z", "T (°C)"] and points[1][:5] == ["P1", 400001.012, 4499998.988, 99.512, 26.0]
    assert points[2][3] is None  # no elevation there
    units = book["Units"]
    assert units[1][:4] == ["Saggio", "1", "12", "US"] and units[1][5] == pytest.approx(1.0)
    assert units[0][-1] == "descrizione" and units[1][-1] == "strato di crollo"
    assert book["Stones"][1][2] == pytest.approx(0.04) and book["Unit proxies"][1][3] == 0.25


def test_exports_offer_excel(tmp_path):
    """The exports that wrote a CSV of a table write an Excel workbook too: the references, the
    cameras, the measuring points."""
    from pyarchinit_photogrammetry.engine.probe import write_points

    doc, cm, ct = small_project(tmp_path)
    m = ct.add_marker("target 1", position=np.array([0.0, 0.0, 0.0]), reference=[400001.0, 4499999.0, 99.0])
    m.checkpoint = True
    book = read_xlsx(ct.export_reference(tmp_path / "markers.xlsx"))
    assert book["markers"][0][:4] == ["label", "x_ref", "y_ref", "z_ref"] and book["markers"][1][1] == 400001.0
    assert book["markers"][1][-1] == "check"
    pts = write_points(ct.meta["drawings"]["points"], tmp_path / "punti.xlsx", ct.crs)
    rows = read_xlsx(pts)["punti_misura"]
    assert rows[0][:3] == ["x", "y", "label"] and rows[1][0] == 400001.012


# ---------------------------------------------------------------------------------------- interface
@pytest.fixture
def qapp(tmp_path):
    pytest.importorskip("PySide6")
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    os.environ.setdefault("PPG_RENDERER", "software")
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication

    QSettings.setDefaultFormat(QSettings.Format.IniFormat)
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, str(tmp_path / "settings"))
    return QApplication.instance() or QApplication([])


@pytest.fixture
def menus(qapp, monkeypatch):
    """The context menus shown, as the texts of their entries: QMenu.exec of PySide6 cannot be
    replaced on the class (the menu would wait for a click), so the menus are made of a subclass
    whose exec only records them."""
    from PySide6 import QtWidgets

    from pyarchinit_photogrammetry.gui import table_export

    seen: list[list[str]] = []

    class RecordedMenu(QtWidgets.QMenu):
        def exec(self, *args):
            seen.append([a.text() for a in self.actions() if not a.isSeparator()])
            return None

    monkeypatch.setattr(QtWidgets, "QMenu", RecordedMenu)
    monkeypatch.setattr(table_export, "QMenu", RecordedMenu)
    return seen


def test_table_widget_to_excel(qapp, menus, tmp_path):
    """What a table shows into a workbook: numbers back as numbers (the value under RAW_ROLE when
    a cell keeps one), ticks as booleans, hidden columns left out."""
    from PySide6.QtCore import QPoint, Qt
    from PySide6.QtWidgets import QTableWidget, QTableWidgetItem

    from pyarchinit_photogrammetry.gui import table_export as te

    assert te.parse_cell("1.25") == 1.25 and te.parse_cell("-12") == -12 and te.parse_cell("12.345.678") == 12345678
    assert te.parse_cell("1,5") == 1.5 and te.parse_cell("1.234,56") == pytest.approx(1234.56)  # Italian interface
    assert te.parse_cell("007") == "007" and te.parse_cell("US 12") == "US 12" and te.parse_cell("") is None
    t = QTableWidget(2, 4)
    t.setHorizontalHeaderLabels(["Nome", "Valore", "Usa", "Nascosta"])
    for r, (name, value, raw) in enumerate((("A", "1.2500", None), ("B", "114.835", 114835))):
        t.setItem(r, 0, QTableWidgetItem(name))
        it = QTableWidgetItem(value)
        if raw is not None:
            it.setData(te.RAW_ROLE, raw)
        t.setItem(r, 1, it)
        tick = QTableWidgetItem()
        tick.setFlags(Qt.ItemFlag.ItemIsUserCheckable | Qt.ItemFlag.ItemIsEnabled)
        tick.setCheckState(Qt.CheckState.Checked if r == 0 else Qt.CheckState.Unchecked)
        t.setItem(r, 2, tick)
        t.setItem(r, 3, QTableWidgetItem("x"))
    t.setColumnHidden(3, True)
    headers, rows = te.table_data(t)
    assert headers == ["Nome", "Valore", "Usa"] and rows == [["A", 1.25, True], ["B", 114835, False]]
    out = te.export_table(t, "Prova", str(tmp_path / "prova.xlsx"))
    assert read_xlsx(out)["Prova"] == [["Nome", "Valore", "Usa"], ["A", 1.25, True], ["B", 114835, False]]
    te.install_excel_menu(t, "Prova")
    t.customContextMenuRequested.emit(QPoint(2, 2))
    assert menus == [["Esporta tabella in Excel..."]]


def test_tables_of_the_panes_offer_excel(qapp, menus, tmp_path, monkeypatch):
    """*Esporta tabella in Excel...* in the context menus of the markers, scale bars and cameras of
    the Riferimenti pane (beside their own entries), of the proxies, and for the attributes of a
    layer of the orthomosaic view."""
    from PySide6.QtCore import QPoint

    from pyarchinit_photogrammetry.gui import image_views, panels, table_export

    doc, cm, ct = small_project(tmp_path)
    a = ct.add_marker("target 1", position=np.array([0.0, 0.0, 0.0]), reference=[10.0, 20.0, 1.0])
    b = ct.add_marker("target 2", position=np.array([1.0, 0.0, 0.0]), reference=[11.0, 20.0, 1.0])
    ct.add_scalebar(a, b, 1.0)
    pane = panels.ReferencePane()
    pane.set_chunk(ct)
    seen = menus
    pane._marker_menu(QPoint(2, 2))
    pane._bar_menu(QPoint(2, 2))
    pane.cams.customContextMenuRequested.emit(QPoint(2, 2))
    assert len(seen) == 3 and all("Esporta tabella in Excel..." in texts for texts in seen)
    assert any(t.startswith("Elimina") for t in seen[1])  # the entries the menu had stay
    out = tmp_path / "marker.xlsx"
    monkeypatch.setattr(table_export.QFileDialog, "getSaveFileName", lambda *a, **k: (str(out), ""))
    menu = table_export.QMenu()  # kept: the action lives in it
    table_export.add_excel_action(menu, pane.table, "Marker").trigger()
    rows = read_xlsx(out)["Marker"]
    assert rows[0][:3] == ["Usa", "Marker", "Check"] and rows[1][:6] == [True, "target 1", False, 10.0, 20.0, 1.0]
    proxies = panels.ProxiesPane()
    proxies.set_proxies(ct.meta["proxies"])
    proxies.table.customContextMenuRequested.emit(QPoint(2, 2))
    assert seen[-1] == ["Esporta tabella in Excel..."]
    # the attributes of a layer
    view = image_views.RasterView()
    view.load(ct.data_dir / "multispectral" / "thermal.tif", "thermal")
    view.add_point(400001.0, 4499999.0, label="P1", layer="Punti di misura", attrs={"temp_c": 26.0})
    view.add_point(400001.5, 4499998.5, label="P2", layer="Punti di misura", attrs={"temp_c": 18.0, "z": 99.5})
    layers = panels.LayersPane()
    layers.set_view(view)
    headers, rows = layers.layer_table("Punti di misura")
    assert headers == ["etichetta", "x", "y", "temp_c", "z"] and rows[1] == ["P2", 400001.5, 4499998.5, 18.0, 99.5]
    layers.tree.setCurrentItem(layers.tree.topLevelItem(0))
    layers.tree.customContextMenuRequested.emit(layers.tree.visualItemRect(layers.tree.topLevelItem(0)).center())
    assert "Esporta gli attributi di Punti di misura in Excel..." in seen[-1]
    out = layers.export_layer("Punti di misura", str(tmp_path / "punti.xlsx"))
    assert read_xlsx(out)["Punti di misura"][2] == ["P2", 400001.5, 4499998.5, 18.0, 99.5]


def test_menu_actions_write_the_report_and_the_tables(qapp, tmp_path, monkeypatch):
    """File > Esporta > Rapporto PDF (the dialog answered by a stand-in: every chunk, English) writes
    the PDF and opens it; File > Esporta > Tabelle (Excel) writes the workbook of the active chunk
    with the measuring points of the window; the result tables offer Excel too."""
    from pyarchinit_photogrammetry.gui import main_window as mwmod

    doc, cm, ct = small_project(tmp_path)
    doc.chunk = ct
    w = mwmod.MainWindow(doc, autoconnect=False)
    w.runner.show_dialog = False
    try:
        texts = [a.text() for a in w.export_menu.actions()]
        assert "Rapporto PDF..." in texts and "Tabelle (Excel)..." in texts
        asked = {}

        def ask(parent, title, fields, intro=""):
            asked.update(title=title, fields=[f.name for f in fields])
            return {"title": "Rilievo", "description": "Prova", "chunks": "all", "language": "en"}

        monkeypatch.setattr(mwmod.ParamDialog, "ask", ask)
        pdf = tmp_path / "out" / "rapporto"
        monkeypatch.setattr(mwmod.QFileDialog, "getSaveFileName", lambda *a, **k: (str(pdf), ""))
        opened = []
        monkeypatch.setattr(mwmod.QDesktopServices, "openUrl", lambda url: opened.append(url.toLocalFile()))
        w.generate_pdf_report()
        w.runner.wait()
        assert asked["title"] == "Rapporto PDF" and asked["fields"] == ["title", "description", "chunks", "language"]
        out = pdf.with_suffix(".pdf")
        assert [Path(p) for p in opened] == [out] and out.read_bytes().startswith(b"%PDF")
        marks = outline(out.read_bytes())
        assert {"Chunk: Campo", "Chunk: Volo termico", "Multispectral", "Thermal"} <= set(marks)
        assert "Rilievo" in pdf_text(out.read_bytes())
        xlsx = tmp_path / "out" / "tabelle.xlsx"
        monkeypatch.setattr(mwmod.QFileDialog, "getSaveFileName", lambda *a, **k: (str(xlsx), ""))
        w.export_tables_xlsx()
        w.runner.wait()
        book = read_xlsx(xlsx)
        assert {"Termico", "Punti di misura", "Camere"} <= set(book)  # the language of the window
        assert [r[0] for r in book["Punti di misura"][1:]] == ["P1", "P2"]
        dialog = mwmod.TableDialog("Classi dalle US", ["US", "Punti", "Area (m²)"], [["US 1", 114835, 1.25]], w)
        dialog._save_xlsx(str(tmp_path / "classi.xlsx"))
        assert read_xlsx(tmp_path / "classi.xlsx")["Classi dalle US"][1] == ["US 1", 114835, 1.25]
        dialog.close()
    finally:
        w.setWindowModified(False)
        w.close()


def test_command_line(tmp_path):
    from typer.testing import CliRunner

    from pyarchinit_photogrammetry.cli import app

    doc, cm, ct = small_project(tmp_path)
    runner = CliRunner()
    res = runner.invoke(app, ["report", str(doc.path), "--lang", "en", "--chunk", "Volo termico", "--html",
                              str(tmp_path / "r.html"), "--pdf", str(tmp_path / "r.pdf")])
    assert res.exit_code == 0, res.output
    assert (tmp_path / "r.pdf").read_bytes().startswith(b"%PDF") and "Thermal" in (tmp_path / "r.html").read_text("utf-8")
    assert "Chunk: Campo" not in outline((tmp_path / "r.pdf").read_bytes())
    res = runner.invoke(app, ["report", str(doc.path)])
    assert res.exit_code == 0, res.output
    assert (tmp_path / "campo_report.pdf").exists()
    res = runner.invoke(app, ["report", str(doc.path), "--lang", "fr"])
    assert res.exit_code != 0
    res = runner.invoke(app, ["tables", str(doc.path), str(tmp_path / "t.xlsx"), "--chunk", "Campo", "--lang", "it"])
    assert res.exit_code == 0, res.output
    assert "Indici" in read_xlsx(tmp_path / "t.xlsx")
    # the same from the chunk
    assert "Thermal" in outline(ct.generate_pdf_report(tmp_path / "c.pdf", language="en", title="T").read_bytes())
    assert "Termico" in read_xlsx(ct.export_tables(tmp_path / "c.xlsx", language="it"))
