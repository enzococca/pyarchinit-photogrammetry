"""The tables of the program as Excel workbooks (openpyxl).

:func:`write_xlsx` writes tables - a sheet name mapped to its headers and rows - into one
workbook: numbers stay numbers (shown with the decimals that make sense for them), empty values
stay empty, the header row is bold and frozen with a filter on it, the columns are as wide as
their contents, the sheet names are made valid for Excel (31 characters, no ``[]:*?/\\``) and
unique.

:func:`project_tables` collects the tables of a chunk: the project, the cameras (reference and
estimated positions, errors, tie points), the sensors and their calibration, the markers, the
scale bars, the products, the point classes, the versions, the processing parameters, the
multispectral bands, indices and zones, the thermal measure and the offsets of its frames, the
measuring points, the spot heights, the stratigraphic units drawn, the stones, the US proxies and
the sections of the project - each only when the chunk has it. ``ppg tables`` and *File > Esporta >
Tabelle (Excel)* write them all, one sheet each.
"""
from __future__ import annotations

import logging
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping

import numpy as np

log = logging.getLogger(__name__)

MAX_SHEET_NAME = 31
_INVALID = re.compile(r"[\[\]:*?/\\]")
HEADER_FILL = "1F4E79"


@dataclass
class Sheet:
    """A table: ``headers``, ``rows`` (values: numbers, texts, booleans, dates or None) and,
    when known, the decimals of every column (None: from the values)."""

    headers: list[str]
    rows: list[list]
    decimals: list[int | None] = field(default_factory=list)


def as_sheet(value: Any) -> Sheet:
    """A :class:`Sheet` from a Sheet, a ``(headers, rows)`` pair or a ``{"headers", "rows",
    "decimals"}`` mapping."""
    if isinstance(value, Sheet):
        return value
    if isinstance(value, Mapping):
        return Sheet(list(value.get("headers") or []), [list(r) for r in value.get("rows") or []],
                     list(value.get("decimals") or []))
    headers, rows = value
    return Sheet(list(headers), [list(r) for r in rows])


def sheet_names(names: Iterable[str], empty: str = "Sheet") -> list[str]:
    """Names Excel accepts, in order and unique (case does not tell two sheets apart): the
    characters ``[]:*?/\\`` replaced, no apostrophe at the ends, at most 31 characters, a number
    added to a repeated name."""
    out: list[str] = []
    taken: set[str] = set()
    for raw in names:
        name = _INVALID.sub("_", str(raw or "")).strip().strip("'").strip() or empty
        if name.lower() == "history":  # reserved by Excel
            name += "_"
        name = name[:MAX_SHEET_NAME].rstrip("'")
        base, k = name, 2
        while name.lower() in taken:
            tail = f" ({k})"
            name = base[:MAX_SHEET_NAME - len(tail)].rstrip("'") + tail
            k += 1
        taken.add(name.lower())
        out.append(name)
    return out


def cell_value(v):
    """A value as a cell holds it: numbers (numpy's too) as numbers, NaN and None empty,
    booleans, dates, everything else as text."""
    if v is None:
        return None
    if isinstance(v, (bool, np.bool_)):
        return bool(v)
    if isinstance(v, (int, np.integer)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        f = float(v)
        return f if math.isfinite(f) else None
    if isinstance(v, (datetime, date)):
        return v
    if isinstance(v, (list, tuple)):
        return ", ".join("" if x is None else str(x) for x in v)
    if isinstance(v, dict):
        import json

        return json.dumps(v, ensure_ascii=False, default=str)
    return str(v)


def number_format(values: list, decimals: int | None = None) -> str | None:
    """The number format of a column: ``decimals`` when given, else from its values (whole
    numbers without decimals, thousands grouped from 10000; coordinates and lengths with three
    decimals, small values with four, tiny ones with six)."""
    nums = [v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if not nums:
        return None
    if decimals is not None:
        return "0" if decimals <= 0 else "0." + "0" * int(decimals)
    if all(isinstance(v, int) or float(v).is_integer() for v in nums):
        return "#,##0" if max(abs(v) for v in nums) >= 10000 else "0"
    top = max(abs(float(v)) for v in nums)
    if top >= 1:
        return "0.000"
    if top >= 0.001:
        return "0.0000"
    return "0.000000"


def _shown(v, fmt: str | None) -> str:
    if isinstance(v, bool) or v is None:
        return "" if v is None else str(v)
    if isinstance(v, (int, float)) and fmt:
        d = len(fmt.split(".")[1]) if "." in fmt else 0
        return f"{float(v):,.{d}f}" if "," in fmt else f"{float(v):.{d}f}"
    return str(v)


def write_xlsx(path, tables: Mapping[str, Any], title: str | None = None) -> Path:
    """``tables`` (sheet name -> :class:`Sheet`, ``(headers, rows)`` or a mapping) into the Excel
    workbook ``path``, one sheet each in their order (see the module docstring). Returns the path."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not tables:
        raise ValueError("no tables to write")
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    names = sheet_names(tables.keys())
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "pyarchinit-photogrammetry"
    if title:
        wb.properties.title = str(title)
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=HEADER_FILL)
    head_align = Alignment(wrap_text=True, vertical="center")
    for name, value in zip(names, tables.values()):
        sheet = as_sheet(value)
        ws = wb.create_sheet(name)
        ncols = max([len(sheet.headers)] + [len(r) for r in sheet.rows] + [0])
        headers = [str(h) for h in sheet.headers] + [""] * (ncols - len(sheet.headers))
        rows = [[cell_value(v) for v in r] + [None] * (ncols - len(r)) for r in sheet.rows]
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=ILLEGAL_CHARACTERS_RE.sub("", h))
            cell.font, cell.fill, cell.alignment = bold, fill, head_align
        formats, given = [], []
        for c in range(ncols):
            dec = sheet.decimals[c] if c < len(sheet.decimals) else None
            formats.append(number_format([r[c] for r in rows], dec))
            given.append(dec is not None)
        for ri, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                if isinstance(v, str):
                    v = ILLEGAL_CHARACTERS_RE.sub("", v)
                cell = ws.cell(row=ri, column=c, value=v)
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"  # a text, never a formula
                elif formats[c - 1] and isinstance(v, (int, float)) and not isinstance(v, bool):
                    fmt = formats[c - 1]
                    if isinstance(v, int) and "." in fmt and not given[c - 1]:  # a count among measures
                        fmt = "#,##0" if abs(v) >= 10000 else "0"  # (the values of the project sheet)
                    cell.number_format = fmt
        for c in range(ncols):
            longest = max([len(w) for w in headers[c].split()] + [len(headers[c]) // 2, 4])
            for r in rows[:2000]:
                longest = max(longest, len(_shown(r[c], formats[c])))
            ws.column_dimensions[get_column_letter(c + 1)].width = min(60, longest + 2)
        if ncols:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(1, len(rows) + 1)}"
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    try:
        wb.save(str(tmp))
        tmp.replace(path)
    finally:
        tmp.unlink(missing_ok=True)
    return path


def read_xlsx(path) -> dict[str, list[list]]:
    """The values of every sheet of a workbook, row by row (the header row first)."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    finally:
        wb.close()


# ---------------------------------------------------------------------------------------- the tables of a chunk
TEXT: dict[str, tuple[str, str]] = {  # key: (Italian, English)
    "s_info": ("Progetto", "Project"),
    "s_cameras": ("Camere", "Cameras"),
    "s_sensors": ("Sensori", "Sensors"),
    "s_markers": ("Marker", "Markers"),
    "s_scalebars": ("Barre di scala", "Scale bars"),
    "s_products": ("Prodotti", "Products"),
    "s_classes": ("Classi", "Point classes"),
    "s_versions": ("Versioni", "Versions"),
    "s_parameters": ("Parametri", "Parameters"),
    "s_bands": ("Bande", "Bands"),
    "s_indices": ("Indici", "Indices"),
    "s_zones": ("Zone {0}", "Zones {0}"),
    "s_thermal": ("Termico", "Thermal"),
    "s_offsets": ("Scostamenti termici", "Thermal offsets"),
    "s_points": ("Punti di misura", "Measurement points"),
    "s_quote": ("Quote", "Spot heights"),
    "s_units": ("US", "Units"),
    "s_stones": ("Pietre", "Stones"),
    "s_proxies": ("Proxy US", "Unit proxies"),
    "s_sections": ("Sezioni", "Sections"),
    "field": ("Campo", "Field"),
    "value": ("Valore", "Value"),
    "label": ("Nome", "Label"),
    "enabled": ("Attiva", "Enabled"),
    "aligned": ("Allineata", "Aligned"),
    "sensor": ("Sensore", "Sensor"),
    "x_est": ("X stimata", "X estimated"),
    "y_est": ("Y stimata", "Y estimated"),
    "z_est": ("Z stimata", "Z estimated"),
    "omega": ("Omega (gradi)", "Omega (degrees)"),
    "phi": ("Phi (gradi)", "Phi (degrees)"),
    "kappa": ("Kappa (gradi)", "Kappa (degrees)"),
    "x_ref": ("X rif.", "X reference"),
    "y_ref": ("Y rif.", "Y reference"),
    "z_ref": ("Z rif.", "Z reference"),
    "acc_x": ("Accur. X (m)", "Accuracy X (m)"),
    "acc_y": ("Accur. Y (m)", "Accuracy Y (m)"),
    "acc_z": ("Accur. Z (m)", "Accuracy Z (m)"),
    "accuracy": ("Accuratezza (m)", "Accuracy (m)"),
    "dx": ("dX (m)", "dX (m)"),
    "dy": ("dY (m)", "dY (m)"),
    "dz": ("dZ (m)", "dZ (m)"),
    "error_m": ("Errore (m)", "Error (m)"),
    "error_px": ("Errore (px)", "Error (px)"),
    "tie_points": ("Punti di legame", "Tie points"),
    "sx": ("Precisione X (m)", "Precision X (m)"),
    "sy": ("Precisione Y (m)", "Precision Y (m)"),
    "sz": ("Precisione Z (m)", "Precision Z (m)"),
    "file": ("File", "File"),
    "key": ("Chiave", "Key"),
    "width": ("Larghezza (px)", "Width (px)"),
    "height": ("Altezza (px)", "Height (px)"),
    "model": ("Modello", "Model"),
    "photos": ("Foto", "Photos"),
    "focal_px": ("Focale (px)", "Focal length (px)"),
    "focal_mm": ("Focale (mm)", "Focal length (mm)"),
    "pixel_mm": ("Pixel (mm)", "Pixel (mm)"),
    "fixed": ("Fissa", "Fixed"),
    "code": ("Codice", "Code"),
    "type": ("Tipo", "Type"),
    "used": ("Usato", "Used"),
    "projections": ("Proiezioni", "Projections"),
    "estimated_flags": ("Stimate (bandierine blu)", "Estimated (blue flags)"),
    "bar": ("Barra", "Bar"),
    "marker_a": ("Marker A", "Marker A"),
    "marker_b": ("Marker B", "Marker B"),
    "distance": ("Distanza (m)", "Distance (m)"),
    "estimated": ("Stimata (m)", "Estimated (m)"),
    "control": ("controllo", "control"),
    "check": ("verifica", "check"),
    "product": ("Prodotto", "Product"),
    "size_mb": ("Dimensione (MB)", "Size (MB)"),
    "elements": ("Punti / facce / celle", "Points / faces / cells"),
    "columns": ("Colonne", "Columns"),
    "rows": ("Righe", "Rows"),
    "resolution": ("Risoluzione (m)", "Resolution (m)"),
    "modified": ("Modificato", "Modified"),
    "class": ("Classe", "Class"),
    "points": ("Punti", "Points"),
    "percent": ("%", "%"),
    "version": ("Versione", "Version"),
    "name": ("Nome", "Name"),
    "created": ("Creata", "Created"),
    "active": ("Attiva", "Active"),
    "parameters": ("Parametri", "Parameters"),
    "step": ("Passo", "Step"),
    "parameter": ("Parametro", "Parameter"),
    "band": ("Banda", "Band"),
    "reference_band": ("Riferimento", "Reference"),
    "panel_factor": ("Fattore del pannello", "Panel factor"),
    "registration": ("Registrazione (correlazione mediana)", "Registration (median correlation)"),
    "registered_on": ("Registrata su", "Registered on"),
    "spread": ("Dispersione dei guadagni (log)", "Spread of the gains (log)"),
    "residual": ("Residuo (log)", "Residual (log)"),
    "min": ("Minimo", "Minimum"),
    "max": ("Massimo", "Maximum"),
    "mean": ("Media", "Mean"),
    "median": ("Mediana", "Median"),
    "p2": ("P2", "P2"),
    "p98": ("P98", "P98"),
    "index": ("Indice", "Index"),
    "formula": ("Formula", "Formula"),
    "zone": ("Zona", "Zone"),
    "from": ("Da", "From"),
    "to": ("A", "To"),
    "area_m2": ("Area (m²)", "Area (m²)"),
    "area_share": ("% dell'area", "% of the area"),
    "camera": ("Camera", "Camera"),
    "capture": ("Ripresa", "Capture"),
    "offset_c": ("Scostamento (°C)", "Offset (°C)"),
    "site": ("Sito", "Site"),
    "area": ("Area", "Area"),
    "us": ("US", "US"),
    "layer": ("Layer", "Layer"),
    "perimeter": ("Perimetro (m)", "Perimeter (m)"),
    "x_centroid": ("X baricentro", "X centroid"),
    "y_centroid": ("Y baricentro", "Y centroid"),
    "raster": ("Raster", "Raster"),
    "stone": ("Pietra", "Stone"),
    "diameter": ("Diametro equivalente (m)", "Equivalent diameter (m)"),
    "proxy": ("Proxy", "Proxy"),
    "volume": ("Volume (m³)", "Volume (m³)"),
    "plan_area": ("Pianta (m²)", "Plan area (m²)"),
    "top_min": ("Superficie: quota minima", "Top: lowest"),
    "top_max": ("Superficie: quota massima", "Top: highest"),
    "bottom_min": ("Base: quota minima", "Bottom: lowest"),
    "bottom_max": ("Base: quota massima", "Bottom: highest"),
    "linked": ("Collegato a PyArchInit", "Linked to PyArchInit"),
    "section": ("Sezione", "Section"),
    "kind": ("Tipo", "Kind"),
    "length": ("Lunghezza (m)", "Length (m)"),
    # info rows
    "i_project": ("Progetto", "Project"),
    "i_project_file": ("File del progetto", "Project file"),
    "i_chunk": ("Chunk", "Chunk"),
    "i_crs": ("Sistema di riferimento", "Reference system"),
    "i_scale": ("Scala del chunk", "Chunk scale"),
    "i_photos": ("Foto", "Photos"),
    "i_aligned": ("Foto allineate", "Aligned photos"),
    "i_sensors": ("Sensori", "Sensors"),
    "i_tie_points": ("Punti di legame", "Tie points"),
    "i_reprojection": ("Errore di riproiezione medio (px)", "Mean reprojection error (px)"),
    "i_date": ("Data", "Date"),
    "i_version": ("Versione del programma", "Program version"),
    "i_site": ("Sito", "Site"),
    # thermal rows
    "t_frames": ("Fotogrammi termici", "Thermal frames"),
    "t_decoder": ("Decodifica", "Decoder"),
    "t_decoded": ("Fotogrammi decodificati", "Frames decoded"),
    "t_emissivity": ("Emissivita'", "Emissivity"),
    "t_reflected": ("Temperatura riflessa (°C)", "Reflected temperature (°C)"),
    "t_ambient": ("Temperatura atmosferica (°C)", "Atmospheric temperature (°C)"),
    "t_humidity": ("Umidita' relativa (%)", "Relative humidity (%)"),
    "t_distance": ("Distanza (m)", "Distance (m)"),
    "t_flight_height": ("altezza di volo", "flight height"),
    "t_scale": ("Scala dei valori interi", "Scale of the integer values"),
    "t_offset": ("Offset dei valori interi", "Offset of the integer values"),
    "t_file": ("Ortofoto termica", "Thermal orthomosaic"),
    "t_resolution": ("Risoluzione (m)", "Resolution (m)"),
    "t_min": ("Temperatura minima (°C)", "Minimum temperature (°C)"),
    "t_max": ("Temperatura massima (°C)", "Maximum temperature (°C)"),
    "t_mean": ("Temperatura media (°C)", "Mean temperature (°C)"),
    "t_median": ("Temperatura mediana (°C)", "Median temperature (°C)"),
    "t_p2": ("Temperatura P2 (°C)", "Temperature P2 (°C)"),
    "t_p98": ("Temperatura P98 (°C)", "Temperature P98 (°C)"),
    "t_spread": ("Scostamenti tra i fotogrammi (°C)", "Offsets between the frames (°C)"),
    "t_anomaly": ("Anomalie", "Anomalies"),
}


class Labels:
    """The headers of the tables in one language."""

    def __init__(self, language: str | None = None):
        from .report_content import language_of

        self.language = language_of(language)
        self._i = 0 if self.language == "it" else 1

    def __call__(self, key: str, *args) -> str:
        pair = TEXT.get(key)
        text = key if pair is None else pair[self._i]
        return text.format(*args) if args else text


def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


def _info(chunk, h: Labels, align: dict) -> Sheet:
    from .. import __version__

    doc = chunk.document
    path = getattr(doc, "path", None)
    rows = [[h("i_project"), path.stem if path is not None else ""],
            [h("i_project_file"), str(path) if path is not None else ""]]
    site = (getattr(doc, "meta", None) or {}).get("sito")
    if site:
        rows.append([h("i_site"), str(site)])
    rows += [[h("i_chunk"), chunk.label], [h("i_crs"), chunk.crs], [h("i_scale"), float(chunk.scale)],
             [h("i_photos"), len(chunk.cameras)], [h("i_aligned"), len(chunk.aligned_cameras)],
             [h("i_sensors"), len(chunk.sensors)]]
    if align:
        rows.append([h("i_tie_points"), align.get("tie_points")])
        if align.get("reprojection_error") is not None:
            rows.append([h("i_reprojection"), round(float(align["reprojection_error"]), 4)])
    rows += [[h("i_date"), datetime.now().replace(microsecond=0)], [h("i_version"), __version__]]
    return Sheet([h("field"), h("value")], rows)


def _cameras(chunk, h: Labels, align: dict) -> Sheet | None:
    if not chunk.cameras:
        return None
    from ..core import geometry
    from ..markers.gcp import camera_reference_errors
    from .exports import opk_from_rotation

    try:
        refs = camera_reference_errors(chunk)
    except Exception:  # noqa: BLE001 - an unknown CRS: the references as they were read
        refs = {}
    per_image = (align or {}).get("per_image") or {}
    names = (chunk.meta.get("sfm") or {}).get("names") or {}
    _s, rot, _t = geometry.decompose_similarity(chunk.transform)
    rows = []
    for cam in chunk.cameras:
        est = chunk.to_world(cam.center) if cam.aligned else [None] * 3
        opk = opk_from_rotation(rot @ cam.transform[:3, :3]) if cam.aligned else (None, None, None)
        entry = refs.get(cam.key, {})
        ref = entry.get("reference")
        ref = list(ref) if ref is not None else list(cam.reference.location or [None] * 3)
        acc = list(cam.reference.accuracy or [None] * 3) if cam.reference.location else [None] * 3
        delta = list(entry["delta"]) if "delta" in entry else [None] * 3
        prec = list(cam.meta.get("precision") or [None] * 3)
        tie = per_image.get(names.get(str(cam.key))) if names else None
        rows.append([cam.label, bool(cam.enabled), bool(cam.aligned), chunk.sensors[cam.sensor_key].label
                     if cam.sensor_key in chunk.sensors else cam.sensor_key,
                     *est, *opk, *ref, *acc, *delta, entry.get("error"), tie, *prec, cam.path])
    headers = [h("label"), h("enabled"), h("aligned"), h("sensor"), h("x_est"), h("y_est"), h("z_est"), h("omega"),
               h("phi"), h("kappa"), h("x_ref"), h("y_ref"), h("z_ref"), h("acc_x"), h("acc_y"), h("acc_z"), h("dx"),
               h("dy"), h("dz"), h("error_m"), h("tie_points"), h("sx"), h("sy"), h("sz"), h("file")]
    dec = [None] * 4 + [3, 3, 3, 5, 5, 5, 3, 3, 3, 3, 3, 3, 4, 4, 4, 4, 0, 4, 4, 4, None]
    return Sheet(headers, rows, dec)


PARAM_NAMES = {"SIMPLE_PINHOLE": ["f", "cx", "cy"], "PINHOLE": ["fx", "fy", "cx", "cy"],
               "SIMPLE_RADIAL": ["f", "cx", "cy", "k"], "RADIAL": ["f", "cx", "cy", "k1", "k2"],
               "OPENCV": ["fx", "fy", "cx", "cy", "k1", "k2", "p1", "p2"],
               "FULL_OPENCV": ["fx", "fy", "cx", "cy", "k1", "k2", "p1", "p2", "k3", "k4", "k5", "k6"],
               "SIMPLE_RADIAL_FISHEYE": ["f", "cx", "cy", "k"], "RADIAL_FISHEYE": ["f", "cx", "cy", "k1", "k2"],
               "OPENCV_FISHEYE": ["fx", "fy", "cx", "cy", "k1", "k2", "k3", "k4"],
               "FOV": ["fx", "fy", "cx", "cy", "omega"],
               "THIN_PRISM_FISHEYE": ["fx", "fy", "cx", "cy", "k1", "k2", "p1", "p2", "k3", "k4", "sx1", "sy1"]}


def _sensors(chunk, h: Labels) -> Sheet | None:
    if not chunk.sensors:
        return None
    count: dict[int, int] = {}
    for c in chunk.cameras:
        count[c.sensor_key] = count.get(c.sensor_key, 0) + 1
    order = ["f", "fx", "fy", "cx", "cy", "k", "k1", "k2", "k3", "k4", "k5", "k6", "p1", "p2", "omega", "sx1", "sy1"]
    used = []
    for s in chunk.sensors.values():
        for n in PARAM_NAMES.get(s.model, []):
            if n not in used:
                used.append(n)
    used = [n for n in order if n in used]
    rows = []
    for s in chunk.sensors.values():
        values = dict(zip(PARAM_NAMES.get(s.model, []), s.params))
        rows.append([s.key, s.label, s.width, s.height, s.model, count.get(s.key, 0), float(s.focal_px),
                     s.focal_length_mm, s.pixel_size_mm, bool(s.fixed), *(values.get(n) for n in used)])
    headers = [h("key"), h("label"), h("width"), h("height"), h("model"), h("photos"), h("focal_px"), h("focal_mm"),
               h("pixel_mm"), h("fixed"), *used]
    dec = [0, None, 0, 0, None, 0, 3, 3, 6, None] + [3 if n in ("f", "fx", "fy", "cx", "cy") else 8 for n in used]
    return Sheet(headers, rows, dec)


def _markers(chunk, h: Labels) -> Sheet | None:
    if not chunk.markers:
        return None
    from ..markers.gcp import marker_errors

    errs = {e["label"]: e for e in marker_errors(chunk)}
    precision = (chunk.meta.get("precision") or {}).get("markers", {})
    rows = []
    for m in chunk.markers:
        e = errs.get(m.label, {})
        ref = list(m.reference.location or [None] * 3)
        est = list(chunk.to_world(m.position)) if m.position is not None else [None] * 3
        acc = float(np.mean(m.reference.accuracy)) if m.reference.location and m.reference.accuracy else None
        prec = list(precision.get(str(m.key)) or [None] * 3)
        rows.append([m.label, m.code, h("check") if m.checkpoint else h("control"), bool(m.reference.enabled), *ref, acc,
                     *est, e.get("dx"), e.get("dy"), e.get("dz"), e.get("error"), e.get("projections"),
                     e.get("estimated"), e.get("pixel_error"), *prec])
    headers = [h("label"), h("code"), h("type"), h("used"), h("x_ref"), h("y_ref"), h("z_ref"), h("accuracy"),
               h("x_est"), h("y_est"), h("z_est"), h("dx"), h("dy"), h("dz"), h("error_m"), h("projections"),
               h("estimated_flags"), h("error_px"), h("sx"), h("sy"), h("sz")]
    dec = [None, 0, None, None, 3, 3, 3, 4, 3, 3, 3, 4, 4, 4, 4, 0, 0, 2, 4, 4, 4]
    return Sheet(headers, rows, dec)


def _scalebars(chunk, h: Labels) -> Sheet | None:
    if not chunk.scalebars:
        return None
    from ..markers.gcp import scalebar_errors

    rows = [[e["label"], e["marker_a"], e["marker_b"], e["distance"], e["accuracy"], h(e["type"]), bool(e["enabled"]),
             e.get("estimated"), e.get("error")] for e in scalebar_errors(chunk)]
    return Sheet([h("bar"), h("marker_a"), h("marker_b"), h("distance"), h("accuracy"), h("type"), h("used"),
                  h("estimated"), h("error_m")], rows, [None, None, None, 4, 4, None, None, 4, 4])


def _size(path: Path) -> int:
    if path.is_dir():
        return sum(f.stat().st_size for f in path.rglob("*") if f.is_file())
    return path.stat().st_size


def _products(chunk, h: Labels) -> Sheet | None:
    from .report_content import _ply_counts, _point_count

    rows = []
    for name in chunk.products:
        p = chunk.product_path(name)
        if p is None or not p.exists():
            continue
        elements, cols, nrows, res = None, None, None, None
        try:
            if name in ("point_cloud", "tie_points"):
                elements = _point_count(p)
            elif name == "model" and p.suffix.lower() == ".ply":
                elements = _ply_counts(p).get("face")
            elif p.suffix.lower() in (".tif", ".tiff"):
                import rasterio

                with rasterio.open(p) as src:
                    cols, nrows, res = src.width, src.height, abs(src.res[0])
                    elements = cols * nrows
        except Exception:  # noqa: BLE001
            pass
        rel = str(p.relative_to(chunk.data_dir)) if p.is_relative_to(chunk.data_dir) else str(p)
        rows.append([name, rel, round(_size(p) / 1e6, 3), elements, cols, nrows, res,
                     datetime.fromtimestamp(p.stat().st_mtime).replace(microsecond=0)])
    if not rows:
        return None
    return Sheet([h("product"), h("file"), h("size_mb"), h("elements"), h("columns"), h("rows"), h("resolution"),
                  h("modified")], rows, [None, None, 3, 0, 0, 0, 4, None])


def _classes(chunk, h: Labels) -> Sheet | None:
    counts = chunk.point_classes()
    if not counts:
        return None
    from ..ai.classify import class_names
    from .report_content import class_label

    names = class_names(chunk)
    total = max(sum(counts.values()), 1)
    rows = [[c, class_label(c, names, h.language), n, 100.0 * n / total] for c, n in sorted(counts.items())]
    return Sheet([h("code"), h("class"), h("points"), h("percent")], rows, [0, None, 0, 2])


def _when(text):
    """An ISO date of the versions as a date for the sheet (the text when it is not one)."""
    try:
        return datetime.fromisoformat(str(text))
    except (TypeError, ValueError):
        return text


def _versions(chunk, h: Labels) -> Sheet | None:
    from ..core import versions as vmod

    rows = []
    for product in vmod.PRODUCTS:
        path = chunk.product_path(product)
        if path is not None and path.exists():
            info = vmod.active_info(chunk, product)
            rows.append([product, "", info.get("name"), _when(info.get("created")), True,
                         str(path.relative_to(chunk.data_dir)) if path.is_relative_to(chunk.data_dir) else str(path),
                         chunk.meta.get(vmod.PARAMS[product])])
        for e in vmod.versions(chunk, product):
            rows.append([product, e.get("id"), e.get("name"), _when(e.get("created")), False,
                         (e.get("files") or {}).get(product), e.get("params")])
    if not rows:
        return None
    return Sheet([h("product"), h("version"), h("name"), h("created"), h("active"), h("file"), h("parameters")], rows)


def _parameters(chunk, h: Labels) -> Sheet | None:
    from .report_content import Words, parameter_rows

    rows = parameter_rows(chunk, Words(h.language))
    if not rows:
        return None
    return Sheet([h("step"), h("parameter"), h("value")], [list(r) for r in rows])


def _raster_stats(path: Path, bands=None) -> list[dict]:
    from .report_content import read_raster, stats

    r = read_raster(path, bands)
    return [stats(r.data[i]) for i in range(r.data.shape[0])]


def _bands(chunk, h: Labels) -> Sheet | None:
    from ..engine.multispectral import NON_SPECTRAL

    ms = chunk.meta.get("multispectral") or {}
    spectral = [b for b in (ms.get("bands") or []) if b not in NON_SPECTRAL]
    if not spectral:
        return None
    quality = ms.get("registration_quality") or {}
    factors = (ms.get("panel") or {}).get("factors") or {}
    balance = ms.get("balance") or {}
    refl_stats = {}
    rel = ms.get("reflectance")
    if rel and (chunk.data_dir / rel).exists():
        names = list(ms.get("reflectance_bands") or [])
        for b, st in zip(names, _raster_stats(chunk.data_dir / rel)):
            refl_stats[b] = st
    rows = []
    for b in spectral:
        cc = [q[b]["correlation"] for q in quality.values() if isinstance(q, dict) and isinstance(q.get(b), dict)]
        on = [q[b].get("on") for q in quality.values() if isinstance(q, dict) and isinstance(q.get(b), dict)]
        st = refl_stats.get(b, {})
        rows.append([b, b == ms.get("reference"), factors.get(b), float(np.median(cc)) if cc else None,
                     max(set(on), key=on.count) if on else None, (balance.get("spread") or {}).get(b),
                     (balance.get("residual") or {}).get(b), st.get("min"), st.get("max"), st.get("mean"), st.get("p2"),
                     st.get("p98")])
    return Sheet([h("band"), h("reference_band"), h("panel_factor"), h("registration"), h("registered_on"), h("spread"),
                  h("residual"), h("min"), h("max"), h("mean"), h("p2"), h("p98")], rows,
                 [None, None, 5, 3, None, 4, 4, 4, 4, 4, 4, 4])


def _indices(chunk, h: Labels) -> Sheet | None:
    from ..engine.multispectral import INDICES

    ms = chunk.meta.get("multispectral") or {}
    rows = []
    for name, info in (ms.get("indices") or {}).items():
        path = chunk.data_dir / info.get("file", "")
        if not info.get("file") or not path.exists():
            continue
        st = _raster_stats(path, [1])[0]
        rows.append([name, info.get("formula") or INDICES.get(name, ("", ""))[0], info["file"], st.get("min"),
                     st.get("max"), st.get("mean"), st.get("median"), st.get("p2"), st.get("p98")])
    if not rows:
        return None
    return Sheet([h("index"), h("formula"), h("file"), h("min"), h("max"), h("mean"), h("median"), h("p2"), h("p98")],
                 rows, [None, None, None, 4, 4, 4, 4, 4, 4])


def _zones(chunk, h: Labels) -> dict[str, Sheet]:
    ms = chunk.meta.get("multispectral") or {}
    out = {}
    for name, summary in (ms.get("zoning") or {}).items():
        breaks = list(summary.get("breaks") or [])
        classes = int(summary.get("classes") or len(breaks) + 1)
        areas = {int(k): float(v) for k, v in (summary.get("area_m2") or {}).items()}
        total = sum(areas.values()) or 1.0
        index = (ms.get("indices") or {}).get(name) or {}
        st = {}
        if index.get("file") and (chunk.data_dir / index["file"]).exists():
            st = _raster_stats(chunk.data_dir / index["file"], [1])[0]
        ends = [st.get("min")] + breaks + [st.get("max")]
        rows = []
        for z in range(1, classes + 1):
            rows.append([z, ends[z - 1] if z - 1 < len(ends) else None, ends[z] if z < len(ends) else None,
                         areas.get(z), 100.0 * areas.get(z, 0.0) / total if areas else None])
        out[h("s_zones", name)] = Sheet([h("zone"), h("from"), h("to"), h("area_m2"), h("area_share")], rows,
                                        [0, 4, 4, 2, 2])
    return out


def _thermal(chunk, h: Labels) -> dict[str, Sheet]:
    from ..engine import thermal
    from .report_content import _decoder_of

    th = chunk.meta.get("thermal") or {}
    cams = [c for c in chunk.cameras if thermal.thermal_file(c)]
    path = thermal.thermal_path(chunk)
    if not cams and path is None:
        return {}
    p = thermal.params_of(chunk)
    rows = [[h("t_frames"), len(cams)], [h("t_decoder"), _decoder_of(chunk)], [h("t_decoded"), th.get("decoded")],
            [h("t_emissivity"), p["emissivity"]], [h("t_reflected"), p["reflected"]], [h("t_ambient"), p["ambient"]],
            [h("t_humidity"), p["humidity"]],
            [h("t_distance"), p["distance"] if p.get("distance") is not None else h("t_flight_height")]]
    if p.get("scale") is not None:
        rows += [[h("t_scale"), p["scale"]], [h("t_offset"), p.get("offset")]]
    if path is not None:
        st = _raster_stats(path, [1])[0]
        st.update({k: float(v) for k, v in (th.get("stats") or {}).items() if k in st})
        rows += [[h("t_file"), str(path.relative_to(chunk.data_dir))], [h("t_resolution"), th.get("resolution")]]
        rows += [[h(f"t_{k}"), st.get(k)] for k in ("min", "max", "mean", "median", "p2", "p98")]
    offsets = th.get("offsets") or {}
    if offsets.get("spread") is not None:
        rows.append([h("t_spread"), offsets["spread"]])
    an = th.get("anomaly") or {}
    if an.get("file"):
        rows.append([h("t_anomaly"), an["file"]])
    out = {h("s_thermal"): Sheet([h("field"), h("value")], rows)}
    by_key = {str(c.key): c for c in chunk.cameras}
    off_rows = [[by_key[k].label, by_key[k].meta.get("capture"), float(v)] for k, v in offsets.items()
                if k != "spread" and k in by_key]
    if off_rows:
        out[h("s_offsets")] = Sheet([h("camera"), h("capture"), h("offset_c")], off_rows, [None, None, 3])
    return out


def _points(chunk, h: Labels) -> Sheet | None:
    from .report_content import POINT_DECIMALS, Words, measurement_points, point_columns, point_header

    points = measurement_points(chunk)
    if not points:
        return None
    cols = point_columns(points)
    w = Words(h.language)
    return Sheet([point_header(w, k) for k in cols], [[p.get(k) for k in cols] for p in points],
                 [POINT_DECIMALS.get(k, 4) if k != "label" else None for k in cols])


def _quote(chunk, h: Labels) -> Sheet | None:
    quote = [dict(q) for q in ((chunk.meta.get("drawings") or {}).get("quote") or []) if "x" in q and "y" in q]
    if not quote:
        return None
    keys: list[str] = []
    for q in quote:
        for k in q:
            if not str(k).startswith("_") and k not in keys:
                keys.append(k)
    keys = [k for k in ("x", "y", "z") if k in keys] + [k for k in keys if k not in ("x", "y", "z")]
    return Sheet(keys, [[q.get(k) for k in keys] for q in quote], [3 if k in ("x", "y", "z") else None for k in keys])


def _units(chunk, h: Labels) -> Sheet | None:
    units = (chunk.meta.get("drawings") or {}).get("units") or []
    if not units:
        return None
    from shapely import wkt

    extra: list[str] = []
    for u in units:
        for k in (u.get("attributes") or {}):
            if not str(k).startswith("_") and k not in extra:
                extra.append(k)
    rows = []
    for u in units:
        try:
            g = wkt.loads(u["wkt"])
        except Exception:  # noqa: BLE001
            g = None
        c = g.centroid if g is not None and not g.is_empty else None
        attrs = u.get("attributes") or {}
        rows.append([u.get("sito"), u.get("area"), u.get("us"), u.get("tipo"), u.get("layer"),
                     g.area if g is not None else None, g.length if g is not None else None,
                     c.x if c is not None else None, c.y if c is not None else None, *(attrs.get(k) for k in extra)])
    return Sheet([h("site"), h("area"), h("us"), h("type"), h("layer"), h("area_m2"), h("perimeter"), h("x_centroid"),
                  h("y_centroid"), *extra], rows, [None, None, None, None, None, 3, 3, 3, 3])


def _stones(chunk, h: Labels) -> Sheet | None:
    stones = chunk.meta.get("stones") or {}
    if not stones:
        return None
    from shapely import wkt

    rows = []
    for raster, items in stones.items():
        for i, text in enumerate(items, 1):
            try:
                g = wkt.loads(text)
            except Exception:  # noqa: BLE001
                continue
            c = g.centroid
            rows.append([raster, i, g.area, g.length, 2.0 * math.sqrt(g.area / math.pi), c.x, c.y])
    if not rows:
        return None
    return Sheet([h("raster"), h("stone"), h("area_m2"), h("perimeter"), h("diameter"), h("x_centroid"),
                  h("y_centroid")], rows, [None, 0, 4, 3, 3, 3, 3])


def _proxies(chunk, h: Labels) -> Sheet | None:
    entries = chunk.meta.get("proxies") or []
    if not entries:
        return None

    def pair(v, i):
        try:
            return float(v[i])
        except (TypeError, IndexError, ValueError):
            return None

    rows = [[e.get("name"), e.get("us"), e.get("area"), e.get("volume_m3"), e.get("area_m2"), pair(e.get("z_top"), 0),
             pair(e.get("z_top"), 1), pair(e.get("z_bottom"), 0), pair(e.get("z_bottom"), 1), bool(e.get("id_media"))]
            for e in entries]
    return Sheet([h("proxy"), h("us"), h("area"), h("volume"), h("plan_area"), h("top_min"), h("top_max"),
                  h("bottom_min"), h("bottom_max"), h("linked")], rows, [None, None, None, 4, 3, 3, 3, 3, 3, None])


def _sections(chunk, h: Labels) -> Sheet | None:
    """The sections of the project (``sezioni/sezioni.json`` beside the project file)."""
    import json

    path = getattr(chunk.document, "path", None)
    f = path.parent / "sezioni" / "sezioni.json" if path is not None else None
    if f is None or not f.exists():
        return None
    try:
        sections = json.loads(f.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    rows = []
    for s in sections:
        p1, p2 = list(s.get("p1") or [None, None]), list(s.get("p2") or [None, None])
        rows.append([s.get("label"), s.get("kind", "vertical"), s.get("length"), _num(p1[0]), _num(p1[1]), _num(p2[0]),
                     _num(p2[1]), s.get("svg"), s.get("dxf")])
    if not rows:
        return None
    return Sheet([h("section"), h("kind"), h("length"), "X1", "Y1", "X2", "Y2", "SVG", "DXF"], rows,
                 [None, None, 3, 3, 3, 3, 3, None, None])


def project_tables(chunk, language: str | None = None, progress=None) -> dict[str, Sheet]:
    """Every table of ``chunk`` (see the module docstring), sheet name -> :class:`Sheet`, in the
    language of the headers (``it`` or ``en``; default: of the interface, else of the system).
    A table that cannot be read is left out (written in the log)."""
    from ..core.progress import as_progress
    from .report_content import alignment_stats

    progress = as_progress(progress)
    h = Labels(language)
    attached = chunk.document is not None
    try:
        align = alignment_stats(chunk) if attached else {}
    except Exception as exc:  # noqa: BLE001
        log.warning("tables: alignment of %s not read: %s", chunk.label, exc)
        align = {}
    builders = [("s_info", lambda: _info(chunk, h, align)), ("s_cameras", lambda: _cameras(chunk, h, align)),
                ("s_sensors", lambda: _sensors(chunk, h)), ("s_markers", lambda: _markers(chunk, h)),
                ("s_scalebars", lambda: _scalebars(chunk, h)), ("s_parameters", lambda: _parameters(chunk, h))]
    if attached:
        builders += [("s_products", lambda: _products(chunk, h)), ("s_classes", lambda: _classes(chunk, h)),
                     ("s_versions", lambda: _versions(chunk, h)), ("s_bands", lambda: _bands(chunk, h)),
                     ("s_indices", lambda: _indices(chunk, h)), ("s_zones", lambda: _zones(chunk, h)),
                     ("s_thermal", lambda: _thermal(chunk, h))]
    builders += [("s_points", lambda: _points(chunk, h)), ("s_quote", lambda: _quote(chunk, h)),
                 ("s_units", lambda: _units(chunk, h)), ("s_stones", lambda: _stones(chunk, h)),
                 ("s_proxies", lambda: _proxies(chunk, h)), ("s_sections", lambda: _sections(chunk, h))]
    out: dict[str, Sheet] = {}
    for i, (key, build) in enumerate(builders):
        progress(i / len(builders), f"Tables of {chunk.label}: {key[2:]}")
        try:
            got = build()
        except Exception as exc:  # noqa: BLE001
            log.warning("tables: %s of %s not written: %s", key, chunk.label, exc, exc_info=True)
            continue
        if isinstance(got, dict):
            out.update(got)
        elif got is not None:
            out[h(key)] = got
    progress(1.0, f"Tables of {chunk.label}: {len(out)}")
    return out


def export_tables(chunk, path, language: str | None = None, progress=None) -> Path:
    """Every table of ``chunk`` into the Excel workbook ``path`` (:func:`project_tables`)."""
    tables = project_tables(chunk, language=language, progress=progress)
    doc = chunk.document
    stem = doc.path.stem if doc is not None and doc.path is not None else ""
    return write_xlsx(path, tables, title=f"{stem} - {chunk.label}" if stem else chunk.label)


__all__ = ["Sheet", "write_xlsx", "read_xlsx", "project_tables", "export_tables", "sheet_names", "cell_value"]
